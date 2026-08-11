#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
phase 0-5 填充计划生成器：读预绑定蓝图 + extracted_data.json，确定性生成 fill_plan_phaseN.json。

背景：phase 0-2 蓝图已固化 table_index/row/col/data_dep，属纯机械映射——这活由脚本秒级完成，
不该让主agent读几MB的 extracted_data 再逐条手写JSON（实测极慢）。
phase 5（第五章合规表15~22）行数不固定（表15 实测33行、表22 实测50行），因此新增
table_rowset 类型：从数组/分组字典/嵌套结构自动展开成行，并生成插行/追加行/删多余行/跨列合并。

分工：
  - table_fill  （固定 rows_map）→ 脚本按 R{r}C{c} 生成 cells
  - table_rowset（动态行集）    → 脚本按数据条数生成 cells + insert_rows/append_rows + delete_rows + merge_cells
  - auto_paragraphs（固定句式段）→ 脚本把蓝图 text_template 渲染成 fill_plan 的 paragraphs
                                  （违法违规/信用/执业资格等只换变量的模板段，
                                    不再给子agent重新创作——业务标准话术直接直出，快且稳）
  - 其余段落/公式类条目          → 脚本不硬猜，输出到 todo 清单，由主agent或按章子agent处理

rows_map 单元格四种写法（table_fill）：
  {"data_dep": "a.b.c"}                        单字段直取
  {"value": "固定文字"}                         固定文字；写 "" 表示**清空该格**
                                               （表1「项目总体情况」「子项目N」是分节标题行，
                                                 value 列必须留空，不得把摘要表数据填进去）
  {"formula": "..."}                            公式行 → 进 todo 交主agent计算
  {"template": "...", "vars": {...}}            **复合单元格**：一格由多个字段拼成
      例（表1「建设内容和规模」= 用地面积+计费机柜+机柜总功率+建设规模+底层资产）：
        {"template": "[[用地面积{land_area}平方米；]][[计费机柜{rack}个；]]{asset}",
         "vars": {"land_area": "sub_projects[0].land_area",
                  "rack": {"dep": "sub_projects[0].billing_rack_count", "format": "num"},
                  "asset": "sub_projects[0].underlying_asset"}}
      [[...]] 为可选段：段内字段有任一为空则整段丢弃（不会留下"用地面积平方米"这种残句）；
      全部 vars 都取不到 → 该格写占位符并进 todo；部分缺失 → 正常出文本 + todo 提示缺哪几个。

用法:
  python gen_phase_fill_plan.py --blueprint <templates/phase0_blueprints.json> \
      --extracted <work_dir>/extracted_data.json \
      --output <work_dir>/fill_plan_phase0.json \
      [--base-vars-out <work_dir>/base_vars.json]   # 仅phase0时生成 base_vars

输出:
  fill_plan_phaseN.json        （tables 部分，可直接喂给 fill_docx.py）
  fill_plan_phaseN.todo.json   （需主agent/子agent处理的条目：段落撰写/公式/未解析字段清单）
  fill_plan_phaseN.phase6.json （仅当蓝图含 apply_stage=phase6 的条目，如表16重建；须最后应用）

来源标注（citations）:
  本项目要求最终 docx 中所有AI填充内容在正文标明来源。脚本从 extracted_data 取值时
  同步取该字段所在对象（或最近祖先对象）的 _attachment_no/_doc_name/_page/**_section**，生成结构化
  citation 写入 fill_plan，由 fill_docx.py 按 templates/citation_rules.json 渲染成括注。
  （_section 用于**无页码材料**：docx/xlsx 如天眼查专业版信用报告，填报告内小节名
   「2.1工商信息」，与 _page 互斥 —— 有它就不必编页码。）
  落位（可在蓝图 target 用 citation_placement 覆盖）：
    - 窄表（cols≤6，如摘要表/表1/表3~表10）→ table_note：汇总到 insert_paragraphs 的表下注
      （逐格加括注必破版）
    - 宽表（有备注列，如表15/16/19/20/22）→ remark_col：短式括注拼进 citation_col 指定列
    - none → 不标注
  溯源字段缺失（有值但无 _attachment_no/_source）的字段会汇总进 .todo.json 的
  missing_citation 清单，提示向上游提取方补溯源——不静默生成"无来源"的内容。
"""

import argparse
import json
import os
import re
import sys

from handoff_gate import HandoffGateError, assert_handoff_ready

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

CELL_KEY = re.compile(r"^R(\d+)C(\d+)$")
PATH_TOKEN = re.compile(r"([^.\[\]]+)|\[(\d+)\]")

# ======================= 来源标注（citations）=======================
# 溯源字段按 extracted_data_schema.json 约定与数据字段同层存放；
# 取不到时向上找最近祖先对象（如 revenue 的溯源常记在 financials 或 originators[0] 上）。
# `_section`：无页码材料（docx/xlsx，如天眼查专业版信用报告）用它替代 `_page` 定位，
# 例如「2.1工商信息」。它在 citation_rules 的 material 话术里是 [[...]] 可选段，
# 取不到就整段丢弃，不会留残渣 —— 因此绝不需要为了"看起来完整"而编页码。
PROV_KEYS = ("_attachment_no", "_doc_name", "_page", "_section", "_source")
NARROW_TABLE_MAX_COLS = 6          # ≤此列数视为窄表 → 表下注（与 citation_rules.json 一致）
FILE_NO_PAT = re.compile(r"^(\d+[-\d]*)")


def _prov_of(obj):
    """从一个 dict 中提取溯源字段（去掉下划线前缀）"""
    if not isinstance(obj, dict):
        return {}
    return {k[1:]: obj[k] for k in PROV_KEYS if not is_blank(obj.get(k))}


def resolve_prov(data, path):
    """解析 data_dep 路径，同时返回溯源信息。

    返回 (found, value, prov)：
      prov = {attachment_no?, doc_name?, page?, source?}
    查找顺序：①值本身是 dict 时先看它；②叶子的父对象；③逐级向上找最近的祖先对象。
    """
    cur = data
    chain = []                      # 沿途经过的容器（从外到内）
    for m in PATH_TOKEN.finditer(path):
        key, idx = m.group(1), m.group(2)
        chain.append(cur)
        if idx is not None:
            i = int(idx)
            if not isinstance(cur, list) or i >= len(cur):
                return False, None, {}
            cur = cur[i]
        else:
            if not isinstance(cur, dict) or key not in cur:
                return False, None, {}
            cur = cur[key]
    prov = _prov_of(cur)
    if not prov:
        for anc in reversed(chain):
            prov = _prov_of(anc)
            if prov:
                break
    return True, cur, prov


def resolve(data, path):
    """兼容旧签名：只返回 (found, value)"""
    found, val, _ = resolve_prov(data, path)
    return found, val


def source_to_attachment(source):
    """_source 是磁盘/图片路径，不能进正文。尽力从中还原「材料编号 + 文件名称」。

    例：'images/13-1-5-1 企业投资项目备案通知书/page_001.png'
        → ('13-1-5-1', '企业投资项目备案通知书')
    取路径中第一个「以编号前缀开头」的组件（与 proofs_index.json 的材料编号口径一致）。
    """
    if is_blank(source):
        return None, None
    parts = [p.strip() for p in str(source).replace("\\", "/").split("/") if p.strip()]
    for comp in reversed(parts):
        stem = os.path.splitext(comp)[0]
        m = FILE_NO_PAT.match(stem)
        if m:
            no = m.group(1).rstrip("-")
            name = stem[len(m.group(1)):].strip(" _-—") or None
            return no, name
    return None, None


def make_citation(prov, field_path=None):
    """把 prov 转成 fill_plan 的 citation 对象。

    优先用 _attachment_no/_doc_name（schema 规定必录、可直接进正文）；
    只有 _source 时从路径还原编号与文件名（兼容尚未补录新字段的历史数据）；
    都取不到 → pending，让「无来源」在产出中显式可见，不静默省略。
    """
    prov = prov or {}
    no = prov.get("attachment_no")
    name = prov.get("doc_name")
    if is_blank(no):
        no2, name2 = source_to_attachment(prov.get("source"))
        no = no or no2
        if is_blank(name):
            name = name2
    if is_blank(no):
        return {"type": "pending", "field": field_path or ""}
    cit = {"type": "material", "attachment_no": str(no)}
    if not is_blank(name):
        cit["doc_name"] = str(name)
    pg = prov.get("page")
    if not is_blank(pg):
        s = str(pg).strip()
        if re.fullmatch(r"\d+(?:-\d+)?", s):
            cit["page"] = s
    # 无页码材料（docx/xlsx）用 _section 定位到报告内小节，如「2.1工商信息」。
    # 只在真的没有 page 时才带，避免「第43页，2.1工商信息」这种双重定位。
    sec = prov.get("section")
    if is_blank(cit.get("page")) and not is_blank(sec):
        cit["section"] = str(sec).strip()
    return cit


def cit_key(cit):
    """去重键：同一份材料同一页（或同一小节）只算一条"""
    if not isinstance(cit, dict):
        return None
    return (cit.get("type"), cit.get("attachment_no"), cit.get("page"),
            cit.get("section"),
            cit.get("field"), cit.get("name"), cit.get("formula"))


class CitationCollector:
    """收集一张表内的全部来源，供表下注汇总 + missing_citation 上报"""

    def __init__(self, entry, target):
        self.no = entry.get("no")
        self.chapter = entry.get("chapter")
        self.table_index = target.get("table_index")
        cols = int(target.get("cols") or 0)
        default = "table_note" if (cols and cols <= NARROW_TABLE_MAX_COLS) else "remark_col"
        self.placement = str(target.get("citation_placement") or default)
        self.citation_col = target.get("citation_col")
        if self.citation_col is not None:
            try:
                self.citation_col = int(self.citation_col)
            except (TypeError, ValueError):
                self.citation_col = None
        # 宽表未指定备注列 → 退化为表下注（否则不知道往哪写）
        if self.placement == "remark_col" and self.citation_col is None:
            self.placement = "table_note"
        self.items = []          # [(key, cit)]
        self.missing = []        # 有值但无溯源的字段路径

    @property
    def enabled(self):
        return self.placement != "none"

    def add(self, prov, field_path, has_value=True):
        if not self.enabled or not has_value:
            return None
        cit = make_citation(prov, field_path)
        if cit.get("type") == "pending":
            if field_path and field_path not in self.missing:
                self.missing.append(field_path)
            return None          # pending 不进表下注（避免整表挂一堆「待填写：来源」）
        k = cit_key(cit)
        if k not in {kk for kk, _ in self.items}:
            self.items.append((k, cit))
        return cit

    def note_item(self):
        """生成 insert_paragraphs 的表下注条目（由 fill_docx 渲染文字）"""
        if self.placement != "table_note" or not self.items or self.table_index is None:
            return None
        return {"after_table_index": self.table_index,
                "citations": [c for _, c in self.items]}

    def todo_item(self):
        if not self.missing:
            return None
        return {"no": self.no, "chapter": self.chapter, "type": "missing_citation",
                "table_index": self.table_index,
                "fields": self.missing,
                "reason": "以下字段有值但 extracted_data 中缺 _attachment_no/_doc_name/_source，"
                          "无法生成正文来源标注。请向上游提取方为这些字段补溯源"
                          "（_attachment_no 取文件名编号前缀、_doc_name 取文件正式名称、_page 取页码）"}



def is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "-", "null"))


# ======================= table_rowset：动态行集展开 =======================

def fmt_amount(v):
    """金额格式化：千分位 + 2位小数；非数值原样返回"""
    try:
        return "{:,.2f}".format(float(v))
    except (TypeError, ValueError):
        return "" if v is None else str(v)


def fmt_num(v):
    try:
        f = float(v)
        return "{:,.2f}".format(f).rstrip("0").rstrip(".") if f % 1 else "{:,}".format(int(f))
    except (TypeError, ValueError):
        return "" if v is None else str(v)


def fmt_by(v, kind):
    """按 rows_map / vars 的 format 声明格式化：amount=千分位2位小数、num=千分位智能位数、其余原样。

    金额行不格式化会写出 "201970.0" 这种口径（标准答案是 "201,970.00"）——
    因此蓝图里所有金额/数量单元格都应显式声明 format。
    """
    if kind == "amount":
        return fmt_amount(v)
    if kind == "num":
        return fmt_num(v)
    return "" if v is None else str(v)


def render_template(tpl, ctx):
    """把 {field} 占位替换为 ctx 中的值（缺失/空 → 空串）。

    支持可选段语法 [[...]]：段内引用的字段只要有一个为空，整段丢弃。
    例："{remark}[[（见附件{attachment_no}）]]" 在 attachment_no 为空时不会留下"（见附件）"。
    """
    s = str(tpl)

    def drop_optional(m):
        seg = m.group(1)
        fields = re.findall(r"\{([A-Za-z0-9_]+)\}", seg)
        if any(is_blank(ctx.get(f)) for f in fields) or not fields:
            return ""
        return seg
    s = re.sub(r"\[\[(.*?)\]\]", drop_optional, s, flags=re.S)

    def sub(m):
        v = ctx.get(m.group(1))
        return "" if is_blank(v) else str(v)
    s = re.sub(r"\{([A-Za-z0-9_]+)\}", sub, s)
    s = re.sub(r"（\s*）|\(\s*\)|【\s*】", "", s)
    # 丢弃首段后，后续段若以分隔符（，、；）开头会残留前导分隔符（如「，2023年1月竣工」）；
    # strip 首尾中英文逗号/顿号/分号，内部分隔符不受影响
    return s.strip("；;、，, ")


def cell_text_spec(data, spec):
    """解析「单元格文本规格」→ (text, got)。

    四种写法（与 rows_map 一致）：
      "固定文字"                                  纯字符串
      {"value": "固定文字"}                        同上（"" 表示清空该格）
      {"data_dep": "a.b.c", "format": "amount"}    单字段直取
      {"template": "...", "vars": {...}}           复合拼接（支持 [[可选段]]）
    取不到值 → ("", False)，调用方决定是留空还是写占位符。
    """
    if spec is None:
        return "", False
    if not isinstance(spec, dict):
        return str(spec), True
    if "value" in spec:
        return str(spec["value"]), True
    if "template" in spec:
        text, _provs, _missing, got = build_composite(data, spec)
        return (text, bool(got))
    dep = spec.get("data_dep")
    if not dep:
        return "", False
    found, val, _prov = resolve_prov(data, dep)
    if not found or is_blank(val) or isinstance(val, (dict, list)):
        return "", False
    return fmt_by(val, spec.get("format")), True


def header_cells_to_cells(target, data, warn=None):
    """把 target.header_cells（{"R0C1": <规格>, ...}）转成 fill_plan 的 cells 列表。

    取不到值的表头格**不写**（保留模版原文）并记入 warn —— 宁可留「第n-3年」也不写空白，
    留着能被第三步「模版表头残留」检查抓到，写空白反而让人以为本来就没表头。
    """
    out = []
    for hk, spec in (target.get("header_cells") or {}).items():
        m = CELL_KEY.match(str(hk))
        if not m:
            continue
        text, got = cell_text_spec(data, spec)
        if not got:
            if warn is not None:
                warn.append({"header_cell": hk, "kind": "表头取值失败，保留模版原文（可能残留「第n-x年」占位）",
                             "spec": spec if not isinstance(spec, dict) else
                             (spec.get("data_dep") or spec.get("template"))})
            continue
        out.append({"row": int(m.group(1)), "col": int(m.group(2)), "text": text})
    return out


def row_ops_from_target(target, data, key):
    """insert_rows / append_rows / delete_rows 的条件门过滤（支持 when_exists/when_missing）"""
    ops = target.get(key) or []
    if key == "append_rows":
        # append_rows 是 [[...]] 或 [{"when_exists":..,"values":[]}]
        out = []
        for op in ops:
            if isinstance(op, dict):
                if cond_ok(data, op):
                    out.append(op.get("values", []))
            else:
                out.append(op)
        return out
    return [{k: v for k, v in op.items() if not k.startswith("when_")}
            for op in ops if isinstance(op, dict) and cond_ok(data, op)]


def col_value(spec, row_ctx, maps, seq):
    """按列规格取值。spec 支持：
       {"from":"field","field":"x","default":"/","format":"amount|num|text"}
       {"from":"const","value":"固定文字"}
       {"from":"template","template":"{remark}（见附件{attachment_no}）"}
       {"from":"map","field":"category","map":"category_seq","default":""}
       {"from":"seq"}                      → 全表流水序号
       {"from":"group_field","field":"stage"}  → 取所属组注入的字段（_group_* 已并入 row_ctx）
    """
    if not isinstance(spec, dict):
        return str(spec)
    kind = spec.get("from", "field")
    fmt = spec.get("format", "text")
    default = spec.get("default", "")
    if kind == "const":
        return str(spec.get("value", ""))
    if kind == "seq":
        return str(seq)
    if kind == "template":
        out = render_template(spec.get("template", ""), row_ctx)
        return out if out else default
    if kind == "map":
        key = row_ctx.get(spec.get("field"))
        table = maps.get(spec.get("map"), {}) or {}
        val = table.get(str(key)) if key is not None else None
        return str(val) if val is not None else default
    # field / group_field
    val = row_ctx.get(spec.get("field"))
    if is_blank(val):
        return default
    if isinstance(val, (dict, list)):
        return default
    if fmt == "amount":
        return fmt_amount(val)
    if fmt == "num":
        return fmt_num(val)
    return str(val)


def row_citation(ctx):
    """从一行数据的上下文取来源。

    合规类记录（表15~表22）本身就带 attachment_no 数据字段（模版硬要求「附件编号必录」），
    优先用它；其次用 _attachment_no/_doc_name/_page/_source 溯源字段。
    """
    prov = _prov_of(ctx)
    for data_key, prov_key in (("attachment_no", "attachment_no"), ("doc_name", "doc_name")):
        if is_blank(prov.get(prov_key)) and not is_blank(ctx.get(data_key)):
            prov[prov_key] = ctx[data_key]
    return make_citation(prov, None)


def flatten_rows(target, data, maps, warn):
    """把数据源展开成行序列 [{"values":[...N列...], "span":[from,to]|None, "merge_text":str}]。

    span 表示该行需要跨列合并（起始列写文字，其余列跳过），用于：
      - nested 的交易环节小标题行（整行跨列）
      - 阶段小计行 / 合计行（标签跨列）
      - span_rules 命中的特殊行（如表15「用地>土地取得方式」行，模版即 gridSpan=5）

    source_kind:
      list         → data_source 是数组，逐条一行
      grouped_list → data_source 是 {分组名: [条目]}，逐组逐条一行（分组名注入 _group）
      nested       → data_source 是 [{...组字段..., items_field: [条目]}]，每组先出跨列小标题行，
                     再逐条出行；可按 subtotal_by 字段变化插入阶段小计行，末尾可加合计行
    """
    columns = target.get("columns", [])
    n_cols = target.get("cols", len(columns))
    kind = target.get("source_kind", "list")
    dep = target.get("data_source")
    found, src = resolve(data, dep) if dep else (False, None)
    if not found or src in (None, {}, []):
        warn.append({"data_source": dep, "kind": "数据源缺失或为空 → 该表保持模版原样并进 todo"})
        return []

    rows = []
    seq = [0]
    span_rules = target.get("span_rules") or []

    def match_span(ctx):
        """span_rules: [{"field":"procedure_name","equals":"土地取得方式","from_col":3,"to_col":7}]"""
        for r in span_rules:
            v = ctx.get(r.get("field"))
            eq = r.get("equals")
            if v is not None and str(v) == str(eq):
                return [int(r.get("from_col", 0)), int(r.get("to_col", n_cols - 1))]
        return None

    def emit(ctx, span=None, merge_text=None, is_special=False):
        if not is_special:
            seq[0] += 1
        if merge_text is not None and span and span[0] == 0 and span[1] >= n_cols - 1:
            vals = [merge_text] + [""] * (n_cols - 1)
            rows.append({"values": vals, "span": span, "merge_text": merge_text, "cit": None})
            return
        vals = [col_value(c, ctx, maps, seq[0]) for c in columns]
        vals += [""] * (n_cols - len(vals))
        vals = vals[:n_cols]
        sp = span if span is not None else match_span(ctx)
        rows.append({"values": vals, "span": sp,
                     "merge_text": merge_text if merge_text is not None
                     else (vals[sp[0]] if sp else None),
                     "cit": row_citation(ctx)})


    def keep(ctx):
        """row_filter：require 字段必须非空；exclude_when 命中则丢弃"""
        rf = target.get("row_filter") or {}
        for f in rf.get("require", []):
            if is_blank(ctx.get(f)):
                return False
        for f, bad in (rf.get("exclude_when") or {}).items():
            if str(ctx.get(f)) in (bad if isinstance(bad, list) else [bad]):
                return False
        return True

    if kind == "list":
        items = src if isinstance(src, list) else []
        for it in items:
            if not isinstance(it, dict):
                continue
            if keep(it):
                emit(it)
    elif kind == "grouped_list":
        groups = src if isinstance(src, dict) else {}
        ctxs = []
        for gname, items in groups.items():
            if str(gname).startswith("$") or not isinstance(items, list):
                continue
            for it in items:
                if not isinstance(it, dict):
                    continue
                ctx = dict(it)
                ctx["_group"] = gname
                if keep(ctx):
                    ctxs.append(ctx)
        # 按官方大类序号稳定排序（保证表15 的 1~9 大类顺序与模版一致，未命中的排最后）
        sort_map, sort_field = target.get("sort_by_map"), target.get("sort_by_field")
        if sort_map and sort_field:
            order = maps.get(sort_map, {}) or {}

            def rank(c):
                try:
                    return int(order.get(str(c.get(sort_field)), 999))
                except (TypeError, ValueError):
                    return 999
            ctxs = sorted(ctxs, key=rank)
        for ctx in ctxs:
            emit(ctx)
    elif kind == "nested":
        items_field = target.get("items_field", "tax_items")
        header_tpl = (target.get("group_header_row") or {}).get("template", "{step_no}.{name}")
        subtotal_by = target.get("subtotal_by")
        sub_cfg = target.get("subtotal_row") or {}
        total_cfg = target.get("total_row") or {}
        groups = src if isinstance(src, list) else []
        cur_scope = None
        scope_sum = {}

        def flush_subtotal(scope):
            if not subtotal_by or scope is None or not sub_cfg:
                return
            label = render_template(sub_cfg.get("label_template", "{scope}阶段小计"),
                                    {"scope": scope})
            vals = [""] * n_cols
            lc = sub_cfg.get("label_cols", [0, 1, 2, 3])
            for c in lc:
                if 0 <= c < n_cols:
                    vals[c] = label
            for c_str, field in (sub_cfg.get("sum_cols") or {}).items():
                c = int(c_str)
                if 0 <= c < n_cols:
                    vals[c] = fmt_amount(scope_sum.get((scope, field), 0))
            for c_str, txt in (sub_cfg.get("fixed_cols") or {}).items():
                c = int(c_str)
                if 0 <= c < n_cols:
                    vals[c] = txt
            rows.append({"values": vals, "span": [min(lc), max(lc)] if lc else None,
                         "merge_text": label, "special": "subtotal"})

        for g in groups:
            if not isinstance(g, dict):
                continue
            scope = g.get(subtotal_by) if subtotal_by else None
            if subtotal_by and cur_scope is not None and scope != cur_scope:
                flush_subtotal(cur_scope)
            cur_scope = scope
            emit(g, span=[0, n_cols - 1], merge_text=render_template(header_tpl, g), is_special=True)
            for it in g.get(items_field, []) or []:
                if not isinstance(it, dict):
                    continue
                ctx = dict(g)
                ctx.pop(items_field, None)
                ctx.update(it)          # 条目字段覆盖组字段
                ctx["_group"] = scope
                if not keep(ctx):
                    continue
                emit(ctx)
                for c_str, field in (sub_cfg.get("sum_cols") or {}).items():
                    try:
                        scope_sum[(scope, field)] = scope_sum.get((scope, field), 0) + float(it.get(field) or 0)
                    except (TypeError, ValueError):
                        pass
        if subtotal_by:
            flush_subtotal(cur_scope)
        if total_cfg:
            vals = [""] * n_cols
            lc = total_cfg.get("label_cols", [0, 1, 2, 3])
            label = total_cfg.get("label", "合计")
            for c in lc:
                if 0 <= c < n_cols:
                    vals[c] = label
            for c_str, dep2 in (total_cfg.get("value_cols") or {}).items():
                c = int(c_str)
                f, v = resolve(data, dep2)
                if 0 <= c < n_cols:
                    vals[c] = fmt_amount(v) if f and not is_blank(v) else "【待填写：合计】"
            for c_str, txt in (total_cfg.get("fixed_cols") or {}).items():
                c = int(c_str)
                if 0 <= c < n_cols:
                    vals[c] = txt
            rows.append({"values": vals, "span": [min(lc), max(lc)] if lc else None,
                         "merge_text": label, "special": "total"})
    else:
        warn.append({"data_source": dep, "kind": "未知 source_kind=%s" % kind})
    return rows


def build_rowset_item(entry, data, todo, notes=None, citation_mode="inline"):
    """生成 table_rowset 类型的 fill_plan tables 条目。

    模版结构参数：
      rows            模版总行数
      header_rows     表头行数（数据从该索引开始写）
      has_total_row   模版最后一行是否为「合计」行（表17/18 有）
      total_row_map   合计行的列填法（label/sum_cols/fixed_cols）
    行数策略：优先覆盖模版已有数据行；不够则插行/追加行；多余则删行。
    来源标注：宽表按 citation_col 把短式括注挂到指定列（通常「备注」列）；
              窄表或未指定 citation_col → 汇总为表下注（insert_paragraphs）。
    """
    target = entry.get("target", {})
    maps = entry.get("maps", {}) or {}
    warn = []
    rows = flatten_rows(target, data, maps, warn)
    col = CitationCollector(entry, target) if citation_mode != "none" else None


    ti = target.get("table_index")
    locate = {"table_index": ti}
    if target.get("header_hint"):
        locate["header_hint"] = target["header_hint"]

    if not rows:
        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "table_rowset_empty", "table_index": ti,
                     "data_source": target.get("data_source"),
                     "reason": "数据源为空，表格保持模版原样；请向上游提取方补数据，或确认本表不涉及",
                     "detail": warn})
        return None

    tpl_rows = int(target.get("rows", 0))
    h = int(target.get("header_rows", 1))
    has_total = bool(target.get("has_total_row"))
    region_end = tpl_rows - 1 if has_total else tpl_rows          # 数据行区间 [h, region_end)
    region = max(region_end - h, 0)
    n = len(rows)

    item = {"locate": locate}
    cells, insert_rows, append_rows, delete_rows, merges = [], [], [], [], []

    # 表头改写（可选）：官方模版有些表头不带单位（如表2「资产规模（长度、面积等）」
    # 「决算总投资（如有）」），不落实单位就无法求合计、也与数据列口径不一致；
    # 第四章表4-1/4-3 更是必须把「第n-3年」这类占位表头换成实际年份（值支持 data_dep）。
    cells.extend(header_cells_to_cells(target, data, warn))

    if n > region:
        extra = n - region
        if has_total:
            # 合计行必须留在最后 → 用 insert_rows 在数据区末尾依次插入
            anchor = h + region - 1
            for k in range(extra):
                insert_rows.append({"after_row": anchor + k, "values": []})
        else:
            for _ in range(extra):
                append_rows.append([])
    elif n < region:
        for r in range(h + n, region_end):
            delete_rows.append({"row": r})

    # 单元格按「最终表」索引写入
    for i, row in enumerate(rows):
        fr = h + i
        span = row.get("span")
        rc = row.get("cit") if col and col.enabled else None
        if rc is not None and rc.get("type") == "pending":
            rc = None           # 无来源不硬塞占位，统一进 missing_citation 清单
        cit_col = col.citation_col if (col and col.placement == "remark_col") else None
        for c, val in enumerate(row["values"]):
            if span and span[0] < c <= span[1]:
                continue      # 跨列区间只写起始列，其余由 merge 覆盖
            attach = (rc if (cit_col is not None and c == cit_col and rc) else None)
            if val == "" and attach is None:
                continue
            cell = {"row": fr, "col": c, "text": val}
            if attach:
                cell["citation"] = attach
            cells.append(cell)
        if rc and col.placement == "table_note":
            k = cit_key(rc)
            if k not in {kk for kk, _ in col.items}:
                col.items.append((k, rc))
        if col and col.enabled and row.get("cit") and row["cit"].get("type") == "pending" \
                and not (span and span[0] == 0 and span[1] >= int(target.get("cols", 8)) - 1):
            col.missing.append("行%d（%s）" % (fr, str(row["values"][:2])))
        if span and span[1] > span[0]:
            merges.append({"row": fr, "from_col": span[0],
                           "to_col": min(span[1], int(target.get("cols", 8)) - 1),
                           "text": row.get("merge_text") or ""})


    # 合计行（模版自带的最后一行）
    if has_total:
        total_row_idx = h + n
        tm = target.get("total_row_map") or {}
        label = tm.get("label", "合计")
        fmt_cols = tm.get("format_cols") or {}      # {"4": "num"}：整数列（如机柜数）别带2位小数
        for c in tm.get("label_cols", [0]):
            cells.append({"row": total_row_idx, "col": c, "text": label})
        for c_str, field in (tm.get("sum_cols") or {}).items():
            total = 0.0
            got = False
            for row in rows:
                try:
                    total += float(str(row["values"][int(c_str)]).replace(",", ""))
                    got = True
                except (TypeError, ValueError, IndexError):
                    pass
            cells.append({"row": total_row_idx, "col": int(c_str),
                          "text": fmt_by(total, fmt_cols.get(str(c_str), "amount")) if got
                          else "【待填写：合计%s】" % field})
        for c_str, txt in (tm.get("fixed_cols") or {}).items():
            cells.append({"row": total_row_idx, "col": int(c_str), "text": txt})

    if cells:
        item["cells"] = cells
    if insert_rows:
        item["insert_rows"] = insert_rows
    if append_rows:
        item["append_rows"] = append_rows
    if delete_rows:
        item["delete_rows"] = delete_rows
    if merges:
        item["merge_cells"] = merges

    n_ph = sum(1 for c in cells if "【待填写" in c["text"])
    n_cit = sum(1 for c in cells if c.get("citation"))
    if col:
        note = col.note_item()
        if note is not None and notes is not None:
            notes.append(note)
            n_cit += len(note["citations"])
        td = col.todo_item()
        if td:
            todo.append(td)
    todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                 "type": "table_rowset_report", "table_index": ti,
                 "reason": "行集已确定性生成，仅供核对：模版%d行(表头%d/数据区%d%s) → 实际数据%d行"
                           % (tpl_rows, h, region, "+合计1" if has_total else "", n),
                 "generated_rows": n, "placeholders": n_ph,
                 "citations": n_cit,
                 "citation_placement": (col.placement if col else "none"),
                 "warnings": warn})
    return item


def build_composite(data, spec):
    """复合单元格：按 spec["template"] + spec["vars"] 把多个 data_dep 拼成一格文本。

    vars 的值有两种写法：
      "a.b.c"                                   直接给路径
      {"dep": "a.b.c", "format": "amount|num|text", "suffix": "万元"}
    返回 (text, provs, missing, got)：
      provs   [(dep, prov)]  已取到值的字段溯源，供表下注汇总
      missing [dep]          取不到值的字段路径（进 todo 提示回补）
      got     是否至少取到一个字段（全空时调用方写占位符）
    """
    ctx, provs, missing = {}, [], []
    for var, vspec in (spec.get("vars") or {}).items():
        if isinstance(vspec, dict):
            vdep, vfmt, suffix = vspec.get("dep"), vspec.get("format", "text"), vspec.get("suffix", "")
        else:
            vdep, vfmt, suffix = vspec, "text", ""
        if not vdep:
            continue
        found, val, prov = resolve_prov(data, vdep)
        if not found or is_blank(val) or isinstance(val, (dict, list)):
            missing.append(vdep)
            continue
        ctx[var] = "%s%s" % (fmt_by(val, vfmt), suffix)
        provs.append((vdep, prov))
    text = render_template(spec.get("template", ""), ctx)
    return text, provs, missing, bool(ctx and text)


def path_present(data, path):
    """条件判定：路径存在且非空（用于 when_exists / when_missing）"""
    if not path:
        return True
    found, val, _ = resolve_prov(data, path)
    if not found or val is None:
        return False
    if isinstance(val, (list, dict)):
        return bool([k for k in val if not str(k).startswith("$")]) if isinstance(val, dict) else bool(val)
    return not is_blank(val)


def cond_ok(data, spec):
    """rows_map 条目 / delete_rows 条目的条件门：
       when_exists  = "sub_projects[1]"  → 该路径有值才生效
       when_missing = "sub_projects[1]"  → 该路径为空才生效
    用途：首发项目只有 1 个子项目，表1 的「子项目2」标题行既不该填、还要连同占位行一起删。
    """
    we = spec.get("when_exists")
    wm = spec.get("when_missing")
    if we and not path_present(data, we):
        return False
    if wm and path_present(data, wm):
        return False
    return True


# ======================= 可模板化段落直出（auto_paragraphs） =======================
# 背景：phase2 的违法违规/信用/执业资格等小节是业务标准话术——整段固定，只换主体名/
# 查询日/附件号几个变量（业务方 AI test 范式：「下方文字中替换【】内容即可输出，
# 其余无需变动」）。之前这些段进 todo 再转子agent创作，既慢又会把固定话术写走样；
# 现在由脚本按蓝图模板确定性渲染，直接进 fill_plan 的 paragraphs。
#
# 蓝图 target.auto_paragraphs 规格：
#   {"foreach": "entities.project_companies",        # 数组路径：逐主体渲染
#    "foreach_paths": ["entities.law_firm", ...],    # 或：固定对象路径清单（中介机构）
#    "sections": [                                    # 两者都不给 = 单次渲染（实体=全局 data）
#      {"match": "（1）说明项目公司近3年在投资建设",  # 官方模版段落锚点（全文唯一子串）
#       "heading": "（1）违法违规情况",             # 可选：替换后的首行小标题
#       "template_key": "text_template_violation",  # 取 target[该key] 作模板
#       "text": "不涉及。",                        # 或：直接给文本（与 template_key 二选一）
#       "requires": "compliance_credit"}]}           # 可选：实体缺该结构时计入缺口报告
#
# 模板变量写法：{entity_name}=主体名（short_name 优先）；{a.b.c}=相对当前主体取值，
# 取不到再按全局 data 取；{a.b.c 顿号连接}=数组用「、」连接。缺字段不编造——
# 原位写【待填写：字段路径】并汇总进 todo（missing_template_fields）。
# 来源标注：这类段落的附件号已内嵌在话术里（「详见附件X」），不另加括注。
AUTO_PARA_VAR = re.compile(r"\{([A-Za-z0-9_\.\[\]']+)( 顿号连接)?\}")

# 提取中间态里的泛称垃圾条目（上游把原文「本公司」当主体名提了一条）——不参与渲染
GENERIC_ENTITY_NAMES = {"本公司", "该公司", "我公司", "贵公司", "公司", "发起人", "原始权益人", "项目公司"}


def merge_entity_copies(arr):
    """同一主体的多个提取副本按名称收敛为一条（字段取先见非空）。

    entities.originators 等数组是提取中间态：同一主体被每份材料各提一次
    （实测 13 条里只有 2 个真主体），不收敛会把同一主体的模板段重复直出 N 次。
    无名/泛称条目剔除；名称按去空白精确匹配（不做模糊归一，全称/简称共存时各留一条，
    由上游提取方保证同主体同名）。
    """
    merged, order = {}, []
    for e in arr:
        if not isinstance(e, dict):
            continue
        name = str(e.get("name") or "").strip()
        if not name or name in GENERIC_ENTITY_NAMES:
            continue
        key = re.sub(r"[\s　]+", "", name)
        if key not in merged:
            merged[key] = {}
            order.append(key)
        tgt = merged[key]
        for k, v in e.items():
            if (k not in tgt or is_blank(tgt[k])) and not is_blank(v):
                tgt[k] = v
    return [merged[k] for k in order]


def entity_display_name(e):
    """主体展示名：short_name 优先，但泛称（如提取方把「本公司」当简称）降级用全称"""
    sn = str(e.get("short_name") or "").strip()
    if sn and sn.lower() != "none" and sn not in GENERIC_ENTITY_NAMES:
        return sn
    return e.get("name")


def select_entities(arr, spec, base_vars):
    """副本收敛后按 base_vars 权威名单 / role 过滤，选出真正的渲染主体。

    entities.originators/project_companies 除副本外还混入无关实体（实测 originators
    13 条含评估机构/财顾，project_companies 26 条含 20+ 兄弟项目公司）：
    1) spec.authority_vars 指定 base_vars 权威键（如 ["originator"]）→ 按名单选取，
       同一白名单名命中的全称/简称副本合并；名单名在 extracted 找不到时也造空实体
       （段落不丢，字段全【待填写】）。
    2) 无 base_vars 时退回 spec.role_filter（按 role 含关键词过滤）。
    3) 都没有 → 返回收敛后全部条目。
    """
    ents = merge_entity_copies(arr)
    wl = []
    for k in spec.get("authority_vars") or []:
        v = (base_vars or {}).get(k)
        if isinstance(v, list):
            # base_vars 数组键（如 originators 权威名单）：逐个进白名单
            wl.extend(str(x).strip() for x in v if not is_blank(x))
        elif not is_blank(v):
            wl.append(str(v).strip())
    if wl:
        out = []
        for w in wl:
            wn = re.sub(r"[\s　]+", "", w)
            hit = {}
            for e in ents:
                en = re.sub(r"[\s　]+", "", str(e.get("name") or ""))
                if en and (en in wn or wn in en):
                    for fk, fv in e.items():
                        if (fk not in hit or is_blank(hit[fk])) and not is_blank(fv):
                            hit[fk] = fv
            hit["name"] = w   # 权威全称为准
            out.append(hit)
        return out
    rf = spec.get("role_filter")
    if rf:
        filtered = [e for e in ents if rf in str(e.get("role") or "")]
        if filtered:
            return filtered
    return ents


def render_para_template(tpl, entity, data, entity_name, missing):
    """渲染一个主体的一段模板文本（缺字段原位【待填写】，不丢段不编造）"""
    def rv(path):
        if path == "entity_name":
            return entity_name
        if isinstance(entity, dict):
            found, val, _ = resolve_prov(entity, path)
            if found:
                return val
        found, val, _ = resolve_prov(data, path)
        return val if found else None

    def sub(m):
        path, join = m.group(1), m.group(2)
        v = rv(path)
        if isinstance(v, list):
            v = "、".join(str(x) for x in v if not is_blank(x))
        if is_blank(v) or isinstance(v, dict):
            if path not in missing:
                missing.append(path)
            return "【待填写：%s】" % path
        return str(v)
    return AUTO_PARA_VAR.sub(sub, str(tpl)).strip()


def build_auto_paragraphs(entry, data, todo, base_vars=None):
    """把一个带 auto_paragraphs 的蓝图条目渲染成 fill_plan 的 paragraphs 条目列表"""
    target = entry.get("target", {})
    spec = target.get("auto_paragraphs") or {}
    if not spec.get("sections"):
        return []

    entities = []          # [(entity_dict_or_None, 主体名)]
    if spec.get("foreach"):
        found, arr = resolve(data, spec["foreach"])
        if found and isinstance(arr, list):
            for e in select_entities(arr, spec, base_vars):
                entities.append((e, entity_display_name(e)))
    elif spec.get("foreach_paths"):
        for p in spec["foreach_paths"]:
            found, obj = resolve(data, p)
            if found and isinstance(obj, dict) and not is_blank(obj.get("name")):
                entities.append((obj, entity_display_name(obj)))
    else:
        entities = [(None, None)]

    if not entities:
        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "auto_paragraphs_no_data",
                     "data_source": spec.get("foreach") or spec.get("foreach_paths"),
                     "reason": "模板段数据源为空，未生成段落；请向上游提取方补数据或确认本节不涉及"})
        return []

    paras = []
    for sec in spec["sections"]:
        tpl = sec.get("text")
        if tpl is None:
            tpl = target.get(sec.get("template_key") or "", "")
        if not tpl or not sec.get("match"):
            continue
        lines = [sec["heading"]] if sec.get("heading") else []
        missing, req_missing = [], []
        for ent, name in entities:
            req = sec.get("requires")
            if ent is not None and req:
                found, val, _ = resolve_prov(ent, req)
                if not found or is_blank(val) or (isinstance(val, dict) and not val):
                    req_missing.append(name or "?")
            lines.append(render_para_template(
                tpl, ent, data, name or "【待填写：主体名称】", missing))
        item = {"match": sec["match"], "replace": "\n".join(lines)}
        if sec.get("section"):
            item["section"] = sec["section"]
        paras.append(item)
        report = {"no": entry.get("no"), "chapter": entry.get("chapter"),
                  "type": "auto_paragraphs_report", "match": sec["match"],
                  "entities": len(entities),
                  "reason": "模板段已确定性直出（%d个主体×1段），仅供核对；附件号已内嵌话术，不另加括注" % len(entities)}
        if missing:
            report["type"] = "missing_template_fields"
            report["fields"] = missing
            report["reason"] = ("模板段已直出，但以下字段取不到值——段内已原位写【待填写：字段】。"
                                "请向上游提取方补数据；严禁子agent自行编造日期/网站/附件号回填")
        if req_missing:
            report["entities_missing_struct"] = req_missing
        todo.append(report)
    return paras


# ======================= table_rebuild：结果态表格重建（ch2+ 表格首选出口） =======================
# 「操作态填空」（table_fill/table_rowset 的 R#C# 坐标 + 删插行）是表格错位/返工的总根源。
# table_rebuild 直出 fill_plan 的 rebuild_tables 条目（执行语义见 fill_docx.py）：
# caption 文本锚定位 → 读旧表按字段名合并已填内容 → 按数据整表新建替换。
# 行数/列数/合并由数据决定，vMerge 陷阱、行号位移、table_index 偏移对它全部失效。
# 蓝图 target 结构：
#   {"type": "table_rebuild", "mode": "kv"|"grid",
#    "locate": {"title_keyword": "表5  发起人…", "occurrence": 1},
#    "create_after": "（2）财务状况",           # 模版缺失表：locate 找不到时的建表锚
#    "caption": "表#  发起人（原始权益人{entity_no}）…",  # 编号一律写「表#」；
#                                              # foreach 时支持 {entity_no}/{entity_name}
#    "merge_existing": true,                    # 透传 fill_docx（默认 true=续填合并）
#    "rows": [{"label": "公司名称", "value": {"data_dep": "$entity.name"},
#              "when_exists": "…"}],            # kv；value 支持 data_dep/template/value 三写法
#    "headers": ["项目", {"data_dep": "…", "fallback": "第n-3年"}],  # grid 表头（值可取数，
#                                              #  年份表头从此取实际值，替代 header_cells 补丁）
#    "row_sets": {…},                           # grid 数据行：与 table_rowset 的 flatten 语义一致
#                                              #  （data_source/source_kind/columns/span_rules/
#                                              #    subtotal/total_row…），span 转为 colspan
#    "static_rows": [[<格规格>, …], …],         # grid 备选：固定行模板逐格取数（格规格同 value，
#                                              #  可加 "colspan"/"placeholder"）
#    "citation_col": 7,                         # grid 行级来源短式进指定列；不给则汇总表后注
#    "foreach": {"path": "entities.originators", "authority_vars": ["originator"]}}
#      → 多主体自动展开：第 1 个主体 locate 官方表；第 2 个及以后自动链式 create_after
#        （锚=前一张表的 caption 文本，fill_docx 建表时会越过表格实体与表下注插入）。
#        行内 data_dep 用 "$entity." 前缀取当前主体字段。
# 数据源为空：官方已有表 → 保持模版原样进 todo；模版缺失表 → 骨架表（表头+一行跨列占位）。

CN_DIGITS = {"零": 0, "〇": 0, "一": 1, "两": 2, "二": 2, "三": 3, "四": 4,
             "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
# 「不满/未满/不足三年」这类否定表述必须先认出来，否则会被 _POS_PAT 的「满」误判为达标
_NEG_YEARS_PAT = re.compile(r"未满|不满|不足|不到|不够|少于|低于|尚未")
_POS_YEARS_PAT = re.compile(r"满|超过|逾|以上|不少于|不低于")
_NUM_YEARS_PAT = re.compile(r"(\d+(?:\.\d+)?)\s*(?:年|周年)?\s*(?:(\d+)\s*个?月)?")
_CN_YEARS_PAT = re.compile(r"([零〇一二三四五六七八九十两]+)\s*(?:年|周年)")
_CN_MONTHS_PAT = re.compile(r"([零〇一二三四五六七八九十两]+)\s*个?月")
# ⚠️ 上游把**日期**误填进 operated_years 是最危险的情形（如「2024年3月1日投入使用至今」）：
# _NUM_YEARS_PAT 会把 2024 当成年限 → 2024 ≥ 3 → **判定达标、静默不提示**，正好漏掉本功能
# 要防的事。对策是**先把日期子串剥掉再解析剩余文本**（而不是一见日期就判读不出）——
# 这样「3.75年（截至2026年6月30日）」这类合法写法仍能取到 3.75，而「2024年3月1日投入使用
# 至今」剥完只剩「投入使用至今」→ 没有年限 → 报"无法核验"。再叠一道不合理大数闸门（>100 年，
# 挡住裸年份「2024」）。宁可报"无法核验"，也不能把读不准的值判成达标。
_DATE_STRIP_PAT = re.compile(
    r"\d{4}\s*年(?:\s*\d{1,2}\s*月)?(?:\s*\d{1,2}\s*日)?"      # 2024年 / 2024年3月 / 2024年3月1日
    r"|\d{4}\s*[-/.]\s*\d{1,2}(?:\s*[-/.]\s*\d{1,2})?")        # 2024-03-01 / 2024/3
MAX_PLAUSIBLE_YEARS = 100.0


def cn_to_num(s):
    """中文数字 → 整数（只需覆盖 1~99：「三」「十」「十二」「二十」「二十三」）。读不出返回 None"""
    s = str(s or "").strip()
    if not s:
        return None
    if "十" in s:
        left, _, right = s.partition("十")
        tens = CN_DIGITS.get(left, 1) if left else 1
        ones = CN_DIGITS.get(right, 0) if right else 0
        if left and left not in CN_DIGITS:
            return None
        if right and right not in CN_DIGITS:
            return None
        return tens * 10 + ones
    n = 0
    for ch in s:
        if ch not in CN_DIGITS:
            return None
        n = n * 10 + CN_DIGITS[ch]
    return n


def parse_operated_years(v):
    """解析 `operating_performance.operation.operated_years` → (years, kind)。

    kind 取值：
      num        解析出可比较的年限（3.75 / "3.75年" / "3年9个月" / "已满三年"→3.0）
      short      文字明示**不满**门槛（"不满三年"/"运营不足2年"）→ 一律提示
      satisfied  文字明示已达门槛但给不出具体数字（"已运营满规定年限"）→ 视为达标
      unparsable 有值但读不出年限（含**日期误填**与不合理大数）→ 提示"无法核验"
      missing    字段为空 → 提示"无法核验"

    ⚠️ 宁可报"无法核验"，也不能把读不准的值判成达标 —— 判成达标就是静默漏报，
    而本功能存在的意义正是"不满 3 年不许被静默放过"。
    """
    if v is None or isinstance(v, bool):
        return None, ("missing" if v is None else "unparsable")
    if isinstance(v, (int, float)):
        return (None, "unparsable") if float(v) > MAX_PLAUSIBLE_YEARS else (float(v), "num")
    if isinstance(v, dict):
        # 上游偶尔把值裹成 {"value": 3.75, "_page": 43}；取得到就按值判，取不到宁可说读不出
        for k in ("value", "years", "operated_years"):
            if k in v:
                return parse_operated_years(v.get(k))
        return None, "unparsable"
    if isinstance(v, (list, tuple)):
        return None, "unparsable"
    s = str(v).strip().replace(",", "").replace("，", "")
    if not s or s in ("-", "null", "N/A", "待填写"):
        return None, "missing"
    neg = bool(_NEG_YEARS_PAT.search(s))
    # 先剥日期，再在剩余文本里找年限（否则「2024年3月1日投入使用」会被读成 2024 年）
    s_num = _DATE_STRIP_PAT.sub(" ", s)
    years = None
    m = _NUM_YEARS_PAT.search(s_num)
    if m and m.group(1):
        years = float(m.group(1))
        if m.group(2):
            years += int(m.group(2)) / 12.0
    else:
        my = _CN_YEARS_PAT.search(s_num)
        if my:
            y = cn_to_num(my.group(1))
            if y is not None:
                years = float(y)
                mm = _CN_MONTHS_PAT.search(s)
                if mm:
                    mo = cn_to_num(mm.group(1))
                    if mo:
                        years += mo / 12.0
    if years is not None and years > MAX_PLAUSIBLE_YEARS:
        return None, "unparsable"      # 如「2024」这类裸年份：不合理即不认，改报无法核验
    if years is not None:
        return years, ("short" if neg else "num")
    if neg:
        return None, "short"
    if _POS_YEARS_PAT.search(s):
        return None, "satisfied"
    return None, "unparsable"


def fmt_years(years, raw):
    """提示文里怎么显示年限：原文可读就用原文，否则用算出来的数值"""
    if not is_blank(raw) and isinstance(raw, str):
        return raw.strip()
    if years is None:
        return ""
    txt = ("%.2f" % years).rstrip("0").rstrip(".")
    return txt + "年"


def build_compliance_note(entry, data, todo, notes):
    """合规门槛提示段：门槛不达标时产出一条 insert_paragraphs（黄底【待确认：…】）。

    target 结构（见 phase4_blueprints.json 的 no:40）：
      check           目前只支持 "operated_years_min"
      data_source     门槛字段所在对象路径（如 operating_performance.operation）
      years_field     年限字段名（operated_years）
      reason_field    不达标时的说明字段名（under_3_years_reason）
      min_years       门槛年数（3）
      after_paragraph 提示段插在哪个段落之后（默认「1.运营时间」= 官方模版 H3 标题，
                      模版与初稿都保留，比指导文字锚点稳）
      notify_on       哪些判定要插提示（默认 short/unparsable/missing）
      note_templates  各判定的提示文案（支持 {占位} 与 [[可选段]]）

    ⚠️ 只往 notes/todo 里追加，永不返回表格条目、永不影响退出码。
    """
    target = entry.get("target", {}) or {}
    if target.get("check") != "operated_years_min":
        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "compliance_note_unsupported",
                     "reason": "未知的 compliance_note.check=%r，本条目跳过（不影响其余条目）"
                               % target.get("check")})
        return None

    src = target.get("data_source") or "operating_performance.operation"
    years_field = target.get("years_field") or "operated_years"
    reason_field = target.get("reason_field") or "under_3_years_reason"
    try:
        min_years = float(target.get("min_years") or 3)
    except (TypeError, ValueError):
        min_years = 3.0
    notify_on = target.get("notify_on") or ["short", "unparsable", "missing"]
    tpls = target.get("note_templates") or {}

    _, obj, prov = resolve_prov(data, src)
    obj = obj if isinstance(obj, dict) else {}
    raw = obj.get(years_field)
    reason = obj.get(reason_field)
    years, kind = parse_operated_years(raw)

    # 判定：num 且 < 门槛 → short（与文字明示不满同一处置）
    if kind == "num" and years is not None and years < min_years:
        kind = "short"
    ok = (kind == "satisfied") or (kind == "num" and years is not None and years >= min_years)

    ctx = {
        "years": fmt_years(years, raw),
        "raw": "" if raw is None else str(raw),
        "min_years": ("%g" % min_years),
        "reason": "" if is_blank(reason) else str(reason),
        "put_into_use_date": obj.get("put_into_use_date") or "",
        "completion_date": obj.get("completion_date") or "",
        "years_path": "%s.%s" % (src, years_field),
        "reason_path": "%s.%s" % (src, reason_field),
        "doc_name": prov.get("doc_name") or "",
        "attachment_no": prov.get("attachment_no") or "",
    }

    if ok:
        todo.append({
            "no": entry.get("no"), "chapter": entry.get("chapter"),
            "type": "operation_years_ok",
            "operated_years": ctx["years"] or ctx["raw"], "min_years": min_years,
            "reason": "运营年限已达 %g 年门槛（%s），未插入提示段。仍请确认正文（三）1.运营时间"
                      "写的结论与本值一致（写「已运营满三年」而实际不满属带来源括注的错数据）"
                      % (min_years, ctx["years"] or ctx["raw"])})
        return None

    key = kind if kind != "short" else ("short_with_reason" if ctx["reason"] else "short")
    tpl = tpls.get(key) or tpls.get("short" if kind == "short" else kind) or tpls.get("short")
    if is_blank(tpl):
        # 蓝图没给文案也要有提示，兜底一句（绝不静默跳过）
        tpl = ("【待确认：运营时间未通过「运营满{min_years}年」门槛核验（{years_path}="
               "「{raw}」）。请核实年限口径，并在本节说明能够实现长期稳定收益的原因；"
               "核实后删除本提示段。】")
    text = render_template(tpl, ctx)
    if is_blank(text):
        return None

    item = {"text": text, "style": target.get("style")}
    # 幂等：重跑蓝图后在已应用过的基底上再应用一次，不会插出第二段一样的提示
    # （dedupe_key 用判定无关的稳定前缀，如「【待确认：运营时间」，换判定也不会重复插）
    if target.get("skip_if_exists", True):
        item["skip_if_exists"] = True
        if not is_blank(target.get("dedupe_key")):
            item["dedupe_key"] = target["dedupe_key"]
    ap = target.get("after_paragraph")
    ati = target.get("after_table_index")
    if not is_blank(ap):
        item["after_paragraph"] = ap
    elif ati is not None:
        item["after_table_index"] = ati
    else:
        item["after_paragraph"] = "1.运营时间"

    if kind in notify_on and isinstance(notes, list):
        notes.append(item)
        placed = True
    else:
        placed = False

    todo.append({
        "no": entry.get("no"), "chapter": entry.get("chapter"),
        "type": "operation_years_short" if kind == "short" else "operation_years_unverified",
        "kind": kind, "operated_years": ctx["years"] or ctx["raw"] or None,
        "min_years": min_years, "has_reason": bool(ctx["reason"]),
        "anchor": item.get("after_paragraph"), "note_inserted": placed,
        "note_text": text if placed else None,
        "reason": ("已在正文（三）1.运营时间的锚点「%s」后插入黄底【待确认】提示段。"
                   "**只提示不阻断**：本项不会让任何脚本报 FAIL。处置：①核实 %s 的口径"
                   "（投入使用日期/申报基准日）；②确实不满 %g 年的，必须在本节补写"
                   "「能够实现长期稳定收益的原因」（数据源 %s）并与主管部门沟通豁免依据；"
                   "③核实到位后把该提示段删除，不得留到交付稿"
                   % (item.get("after_paragraph"), ctx["years_path"], min_years,
                      ctx["reason_path"])
                  if placed else
                  "判定=%s 未列入 notify_on，本次未插提示段（仅记录）" % kind)})
    return None


def _rebuild_cell_value(data, spec, missing, provs, counters=None):
    """kv 行 value / grid 表头 / static_rows 格的取值：返回 (text, got)，顺带收集溯源与缺口。

    counters（可选 dict）：只有 data_dep/template 这种「真取数」的规格才计入 n_data/n_got，
    纯常量不算——用于判定「整表数据源为空」（全常量的表不存在空不空的问题）。"""
    is_data = isinstance(spec, dict) and (spec.get("data_dep") or "template" in spec)
    if counters is not None and is_data:
        counters["n_data"] = counters.get("n_data", 0) + 1
    if isinstance(spec, dict) and spec.get("data_dep"):
        found, val, prov = resolve_prov(data, spec["data_dep"])
        if not found or is_blank(val) or isinstance(val, (dict, list)):
            missing.append(spec["data_dep"])
            return "", False
        if prov:
            provs.append((spec["data_dep"], prov))
        if counters is not None:
            counters["n_got"] = counters.get("n_got", 0) + 1
        return fmt_by(val, spec.get("format")), True
    if isinstance(spec, dict) and "template" in spec:
        text, pv, miss, got = build_composite(data, spec)
        provs.extend(pv)
        missing.extend(miss)
        if got and counters is not None:
            counters["n_got"] = counters.get("n_got", 0) + 1
        return text, got
    return cell_text_spec(data, spec)


def _fin_num(v):
    """财务金额解析：容忍千分位逗号/全角逗号/空白；解析不了返回 None"""
    if v is None:
        return None
    try:
        return float(str(v).replace(",", "").replace("\uff0c", "").replace(" ", ""))
    except ValueError:
        return None


def make_fin_view(fin, data):
    """财务表动态视图 $fin：按 financials 的实际键推「最近3个会计年度+一期」。

    背景：官方表5 表头是「第n-3年…」占位，旧蓝图把实际年份硬编码在 rows_map
    （$year_note 历史事故：换项目忘改，交付稿残留占位）。这里按数据推导：
      y1/y2/y3      最近3个年度的 financials 子dict（升序、保引用，溯源沿祖先可取）
      q             一期子dict（键形如 2024Q3 / 2024Q3_unaudited / 2026Q1）
      h1..h3        年度表头「2023年/2023年12月31日」
      h4            一期表头「2024年1-9月/2024年9月30日」——月区间优先按申报基准日
                    （project_info.declaration_base_date），解析不了回退用一期键的季度推
      dr1..dr3/drq  资产负债率 = 总负债/总资产*100 保留2位（确定性公式，原口径由主agent
                    计算后补格；**EBITDA 仍一律不算**，蓝图直接写【待填写】常量）
    取不到的键一律缺席 → 蓝图 data_dep 落 fallback/占位并进 missing，不编造。"""
    view = {}
    if not isinstance(fin, dict):
        return view
    # 溯源随视图走（$fin 链上找不到祖先 financials，把溯源字段复制进来）
    for k in ("_attachment_no", "_doc_name", "_page", "_source"):
        if fin.get(k) is not None:
            view[k] = fin[k]
    years = sorted(k for k in fin
                   if re.fullmatch(r"\d{4}", str(k)) and isinstance(fin[k], dict))[-3:]
    for i, yk in enumerate(years, 1):
        view["y%d" % i] = fin[yk]
        view["h%d" % i] = "%s\u5e74/%s\u5e7412\u670831\u65e5" % (yk, yk)
    periods = sorted(k for k in fin
                     if re.match(r"^\d{4}Q\d", str(k)) and isinstance(fin[k], dict))
    if periods:
        qk = periods[-1]
        view["q"] = fin[qk]
        qy = qk[:4]
        found, bd = resolve(data, "project_info.declaration_base_date")
        if not found or is_blank(bd):
            bd = (data.get("declaration_base_date")
                  if isinstance(data, dict) else None)
        m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", str(bd or ""))
        if m and m.group(1) == qy:
            view["h4"] = ("%s\u5e741-%d\u6708/%s\u5e74%d\u6708%d\u65e5"
                          % (qy, int(m.group(2)), qy, int(m.group(2)), int(m.group(3))))
        else:
            qn = re.search(r"Q(\d)", qk)
            if qn:
                mm = int(qn.group(1)) * 3
                dd = 31 if mm in (3, 12) else 30
                view["h4"] = ("%s\u5e741-%d\u6708/%s\u5e74%d\u6708%d\u65e5"
                              % (qy, mm, qy, mm, dd))
    for tag, d in [("dr1", view.get("y1")), ("dr2", view.get("y2")),
                   ("dr3", view.get("y3")), ("drq", view.get("q"))]:
        if not isinstance(d, dict):
            continue
        ta = _fin_num(d.get("total_assets"))
        tl = _fin_num(d.get("total_liabilities"))
        if ta and tl is not None:
            view[tag] = "%.2f" % (tl / ta * 100)
    # 同比增幅（表2-3附表用）：只在两端都是完整年度数据时计算，
    # 旧值为0或任一值缺失该字段不出现（不编造、不用0/None伪装）
    years_txt = [str(y) for y in years]
    if len(years_txt) == 3:
        view["g1_label"] = "%s年较%s年增幅" % (years_txt[1], years_txt[0])
        view["g2_label"] = "%s年较%s年增幅" % (years_txt[2], years_txt[1])
    growth_fields = ("total_assets", "total_liabilities", "revenue", "net_profit",
                     "operating_cash_flow")
    for tag, d_old, d_new in (("g1", view.get("y1"), view.get("y2")),
                              ("g2", view.get("y2"), view.get("y3"))):
        if not isinstance(d_old, dict) or not isinstance(d_new, dict):
            continue
        gd = {}
        for f in growth_fields:
            n, o = _fin_num(d_new.get(f)), _fin_num(d_old.get(f))
            if n is None or o is None or o == 0:
                continue
            gd[f] = "%+.2f%%" % ((n - o) / abs(o) * 100)
        if gd:
            view[tag] = gd
    # 资产负债率变动（百分点，两个百分比数字直接相减，非相对增幅）
    for tag, o_key, n_key in (("dd1", "dr1", "dr2"), ("dd2", "dr2", "dr3")):
        o, n = view.get(o_key), view.get(n_key)
        if o is None or n is None:
            continue
        try:
            view[tag] = "%+.2f个百分点" % (float(n) - float(o))
        except (TypeError, ValueError):
            pass
    return view


def build_rebuild_items(entry, data, todo, citation_mode="inline", base_vars=None):
    """target.type == "table_rebuild" → [rebuild_tables 条目…]（foreach 时一主体一条）。"""
    target = entry.get("target", {})
    mode = target.get("mode", "kv")
    maps = entry.get("maps", {}) or {}
    if target.get("citation_placement") == "none":
        # 条目级关闭来源标注（表22：官方表无备注列，括注/表后注都会破版）
        citation_mode = "none"
    out = []

    fe = target.get("foreach") or {}
    if fe:
        found, arr = resolve(data, fe.get("path", ""))
        ents = select_entities(arr if isinstance(arr, list) else [], fe, base_vars)
        if not ents:
            todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                         "type": "table_rebuild_skipped",
                         "table_name": (target.get("locate") or {}).get("title_keyword")
                                       or target.get("caption"),
                         "data_source": fe.get("path"),
                         "reason": "foreach 主体数组为空 → 本表未生成；补数据后重跑本 phase"})
            return out
    else:
        ents = [None]

    prev_caption = None
    for idx, ent in enumerate(ents):
        dctx = data if ent is None else dict(data, **{"$entity": ent})
        if ent is not None and isinstance(ent.get("financials"), dict):
            dctx["$fin"] = make_fin_view(ent["financials"], data)
        item = {"mode": mode}
        tag = ((target.get("locate") or {}).get("title_keyword")
               or target.get("caption") or "")

        # ---- 定位 / 建表锚 ----
        if idx == 0:
            if target.get("locate"):
                item["locate"] = dict(target["locate"])
            if target.get("create_after"):
                ca = target["create_after"]
                if ent is not None:
                    # 支持 create_after 也带 {entity_no}/{entity_name}：
                    # 当新表需锚在“同一主体的前一张表”时（如同比增幅表锚在
                    # 该主体的财务表后），前表 caption 本身带了主体序号，这里同样无 no_txt（单主体时为空）
                    _no_txt = str(idx + 1) if len(ents) > 1 else ""
                    ca = (str(ca).replace("{entity_no}", _no_txt)
                                 .replace("{entity_name}", str(entity_display_name(ent) or "")))
                item["create_after"] = ca
        else:
            # 多主体副本：链式锚——上一张表的 caption 文本；fill_docx 建表时会越过
            # 锚点段后紧跟的表格实体与表下注，副本天然按主体顺序排在前表之后
            anchor = prev_caption
            if not anchor:
                todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                             "type": "table_rebuild_copy_dropped", "table_name": tag,
                             "reason": "主体%d 副本无链式锚（前一条无 caption/locate），已跳过" % (idx + 1)})
                continue
            item["create_after"] = anchor

        # ---- caption（编号一律「表#」，交付前 renumber_tables.py 统一赋号）----
        cap = target.get("caption")
        if cap and ent is not None:
            # 单主体时 {entity_no} 渲染为空：「原始权益人{entity_no}」→「原始权益人」；
            # 多主体才编号 1/2/3…（副本标题与主体一一对应）
            no_txt = str(idx + 1) if len(ents) > 1 else ""
            cap = (str(cap).replace("{entity_no}", no_txt)
                           .replace("{entity_name}", str(entity_display_name(ent) or "")))
        if idx > 0 and not cap:
            todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                         "type": "table_rebuild_copy_dropped", "table_name": tag,
                         "reason": "多主体副本必须提供 caption 模板（支持{entity_no}/{entity_name}），"
                                   "已跳过主体%d" % (idx + 1)})
            continue
        if cap:
            item["caption"] = cap
            tag = cap
        if target.get("merge_existing") is False:
            item["merge_existing"] = False
        if target.get("merge_options"):
            # 续填合并调参（fill_docx merge_rebuild_grid）：scaffold_headers=模版预印
            # 脚手架列（表15 手续菜单/表22 税种骨架，预印文字不算已填）；key_cols=行键列
            item["merge_options"] = dict(target["merge_options"])
        if target.get("style"):
            item["style"] = target["style"]

        provs, missing, warn, extra_cits = [], [], [], []
        counters = {"n_data": 0, "n_got": 0}
        if mode == "kv":
            rows = []
            for r in (target.get("rows") or []):
                if not isinstance(r, dict) or not cond_ok(dctx, r):
                    continue
                label = str(r.get("label", ""))
                text, got = _rebuild_cell_value(dctx, r.get("value"), missing, provs, counters)
                if not got:
                    text = str(r.get("placeholder") or ("【待填写：%s】" % label))
                rows.append({"label": label, "value": text})
            item["rows"] = rows
            empty = not rows or (counters["n_data"] > 0 and counters["n_got"] == 0)
        else:
            headers = []
            for h in (target.get("headers") or []):
                # 复合表头（表4-4 三行表头）：colspan/rowspan 原样带给 fill_docx，
                # 取值规格剥掉结构键后走同一套 data_dep/template/value 解析
                span, hspec = {}, h
                if isinstance(h, dict):
                    for sk in ("colspan", "rowspan"):
                        sv = int(h.get(sk) or 1)
                        if sv > 1:
                            span[sk] = sv
                    hspec = {k: v for k, v in h.items()
                             if k not in ("colspan", "rowspan", "fallback")}
                text, got = _rebuild_cell_value(dctx, hspec, missing, provs)
                if not got:
                    text = ((h.get("fallback") if isinstance(h, dict) else None)
                            or "【待填写：表头】")
                    warn.append({"header": h if not isinstance(h, dict)
                                 else (h.get("data_dep") or h.get("template")),
                                 "kind": "表头取值失败，已用 fallback/占位"})
                headers.append(dict(span, text=text) if span else text)
            item["headers"] = headers
            # 物理列数 = colspan 合计（row_sets/骨架行按物理列铺格）
            n_cols = max(sum(int(h.get("colspan", 1) or 1) if isinstance(h, dict) else 1
                             for h in headers), 1)
            cit_col = target.get("citation_col")
            grows = []
            # static_rows 恒为前缀行（表4-4 第2/3行复合表头等结构行），
            # row_sets 数据行接其后；二者不再互斥
            for srow in (target.get("static_rows") or []):
                # 行级条件形态：{"when_exists"/"when_missing": 路径, "cells": [...]}
                # ——表3 的「项目公司2/3」行有数据才出行，行数由数据决定
                if isinstance(srow, dict):
                    if not cond_ok(dctx, srow):
                        continue
                    srow = srow.get("cells") or []
                cells_row = []
                for cspec in (srow or []):
                    cs = cspec if isinstance(cspec, dict) else None
                    vspec = ({k: v for k, v in cs.items()
                              if k not in ("colspan", "placeholder")} if cs else cspec)
                    text, got = _rebuild_cell_value(dctx, vspec, missing, provs, counters)
                    if not got and cs is not None:
                        text = str(cs.get("placeholder") or "【待填写】")
                    if cs and cs.get("colspan"):
                        cells_row.append({"text": text, "colspan": int(cs["colspan"])})
                    else:
                        cells_row.append(text)
                grows.append(cells_row)
            n_static = len(grows)
            if target.get("row_sets"):
                rs = dict(target["row_sets"])
                rs.setdefault("cols", n_cols)
                for row in flatten_rows(rs, dctx, maps, warn):
                    rc = row.get("cit")
                    if not rc or rc.get("type") == "pending" or citation_mode == "none":
                        rc = None
                    vals, span = row["values"], row.get("span")
                    cells_row, c = [], 0
                    while c < n_cols:
                        if span and c == int(span[0]) and int(span[1]) > int(span[0]):
                            to_c = min(int(span[1]), n_cols - 1)
                            cells_row.append({"text": row.get("merge_text")
                                              or (vals[c] if c < len(vals) else ""),
                                              "colspan": to_c - c + 1})
                            c = to_c + 1
                            continue
                        v = vals[c] if c < len(vals) else ""
                        if rc is not None and cit_col is not None and c == int(cit_col):
                            cells_row.append({"text": v, "citation": rc})
                            rc = None       # 行级来源已落列，不再进表后注
                        else:
                            cells_row.append(v)
                        c += 1
                    if rc is not None:
                        extra_cits.append(rc)   # 未指定 citation_col → 汇总为表后注
                    grows.append(cells_row)
                # 空判定只看数据行：static 前缀是结构行，不算「有数据」
                empty = len(grows) <= n_static
            else:
                empty = not grows or (counters["n_data"] > 0 and counters["n_got"] == 0)
            item["rows"] = grows

        if empty:
            ds = ((target.get("row_sets") or {}).get("data_source")
                  or (fe.get("path") if fe else None) or "（蓝图字段）")
            if (item.get("locate") or {}).get("title_keyword") and not item.get("create_after"):
                # 官方模版已有这张表：保持模版原样，不产出重建条目
                todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                             "type": "table_rebuild_skipped", "table_name": tag,
                             "data_source": ds,
                             "reason": "数据源为空 → 官方模版表保持原样；请向上游提取方补数据，"
                                       "或确认本表不涉及", "detail": warn})
                continue
            # 模版缺失表：骨架承载（表头 + 一行跨列【待填写】），保住表实体防悬空标题
            ph = "【待填写：本表数据缺失（数据源 %s 为空），请补录后重跑本 phase】" % ds
            if mode == "kv":
                if not item["rows"]:
                    item["rows"] = [{"label": "【待填写】", "value": ph}]
            else:
                # 保留 static 前缀行（复合表头的第2/3行），占位行铺满物理列宽
                item["rows"] = (item.get("rows") or [])[:n_static] + [
                    [{"text": ph, "colspan": n_cols}]]
            todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                         "type": "table_new_placeholder", "table_name": tag,
                         "data_source": ds,
                         "reason": "骨架表（rebuild 承载）：只有表头+一行占位，不是交付形态"})

        # ---- 块级 citations → fill_docx 渲染为表后「注：本表数据…」 ----
        if citation_mode != "none":
            cits, seen = [], set()
            for dep, prov in provs:
                c = make_citation(prov, dep)
                if c.get("type") == "pending":
                    continue
                k = cit_key(c)
                if k not in seen:
                    seen.add(k)
                    cits.append(c)
            for c in extra_cits:
                k = cit_key(c)
                if k not in seen:
                    seen.add(k)
                    cits.append(c)
            if cits:
                item["citations"] = cits

        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "table_rebuild_report", "table": tag, "mode": mode,
                     "reason": "重建条目已确定性生成，仅供核对：%d 行，来源 %d 条%s"
                               % (len(item.get("rows") or []),
                                  len(item.get("citations") or []),
                                  "，取不到值的字段 %d 个" % len(missing) if missing else ""),
                     "missing_fields": missing[:20], "warnings": warn})
        prev_caption = cap or (item.get("locate") or {}).get("title_keyword") or prev_caption
        out.append(item)
    return out


def gen_entry(entry, data, todo, notes=None, citation_mode="inline", paras=None, base_vars=None,
              compliance_notes=True, rebuilds=None):
    """处理一个蓝图条目：表格类返回 fill_plan tables item；模板段直出进 paras；其余记入 todo 返回 None"""
    if entry.get("disabled"):
        # 蓝图里显式停用的条目（如释义表）直接跳过，不进 fill_plan 也不进 todo
        return None
    target = entry.get("target", {})
    if target.get("type") == "compliance_note":
        # 合规门槛提示（运营满3年等）：只产出插入段落 + todo，不产表格条目
        if compliance_notes:
            build_compliance_note(entry, data, todo, notes if notes is not None else [])
        else:
            todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                         "type": "compliance_note_disabled",
                         "reason": "已用 --no-compliance-notes 关闭合规门槛提示段"})
        return None
    if target.get("type") == "table_rebuild":
        items = build_rebuild_items(entry, data, todo, citation_mode, base_vars)
        if rebuilds is not None:
            rebuilds.extend(items)
        return None
    if target.get("type") == "table_rowset":
        return build_rowset_item(entry, data, todo, notes, citation_mode)
    # 固定句式段（违法违规/信用/执业资格等）→ 脚本按模板直出，不转子agent
    auto_done = False
    if target.get("auto_paragraphs"):
        built = build_auto_paragraphs(entry, data, todo, base_vars)
        if paras is not None:
            paras.extend(built)
        auto_done = True
    rows_map = target.get("rows_map")
    if not rows_map or target.get("table_index") is None:
        if auto_done:
            # 模板段已直出（或已记 no_data/missing 报告），不再进「需主agent」todo
            return None
        # 段落/模板/公式主导的条目 → 交给主agent
        todo.append({
            "no": entry.get("no"),
            "chapter": entry.get("chapter"),
            "type": target.get("type", "paragraph/other"),
            "reason": "非表格类（段落撰写/文字模板），需主agent按蓝图与text_templates生成",
            "hint": entry.get("fill_plan_hint"),
        })
        return None

    hint = entry.get("fill_plan_hint", {}) or {}
    locate = hint.get("locate") or {"table_index": target["table_index"]}
    if "header_hint" not in locate and target.get("header_hint"):
        locate["header_hint"] = target["header_hint"]

    # table_fill 是 2~6 列的信息表（摘要表/表1/表3~表10），逐格加括注必破版 →
    # 一律汇总为表下注；蓝图写 citation_placement=none 可关闭
    col = None
    if citation_mode != "none":
        col = CitationCollector(entry, target)
        if col.placement == "remark_col":
            col.placement = "table_note"

    cells = []
    unresolved = []
    hwarn = []
    # 表头改写（第四章表4-2 的年份列表头必须换掉「第n-3年」占位）
    cells.extend(header_cells_to_cells(target, data, hwarn))
    for key, spec in rows_map.items():
        m = CELL_KEY.match(key)
        if not m or not isinstance(spec, dict):
            continue  # 如 "R0": "表头行（…）" 的说明项
        row, col_i = int(m.group(1)), int(m.group(2))
        desc = spec.get("desc", "")
        if not cond_ok(data, spec):
            continue        # when_exists/when_missing 未满足 → 该格不生成（如首发项目的「子项目2」行）
        if "formula" in spec:
            # 公式行交主agent计算（如资产负债率；EBITDA 不走这里——蓝图已改为 value 直接【待填写】，不计算）
            todo.append({"no": entry.get("no"), "cell": key, "type": "formula",
                         "formula": spec["formula"], "desc": desc,
                         "citation_hint": {"type": "computed", "formula": spec["formula"],
                                           "basis": "（填写取数来源，如「2024年审计报告数据」）"},
                         "reason": "公式行需主agent按 extracted_data 计算后补入该单元格；"
                                   "补入时请一并带 citation_hint（computed 类来源标注）"})
            continue
        if "value" in spec:  # 固定文字（写 "" 表示清空该格，如表1的分节标题行）
            cells.append({"row": row, "col": col_i, "text": str(spec["value"])})
            continue
        if "template" in spec:
            # 复合单元格：一格由多个字段拼成（表1 的「建设内容和规模」「开竣工时间」
            # 「项目权属起止时间及剩余年限」都是复合行——只映射单字段会漏掉大半内容）
            text, provs, missing, got = build_composite(data, spec)
            if not got:
                cells.append({"row": row, "col": col_i,
                              "text": "【待填写：%s】" % (desc or key)})
                unresolved.append({"cell": key, "desc": desc, "deps": missing,
                                   "kind": "复合字段的全部来源字段均为空"})
                continue
            cells.append({"row": row, "col": col_i, "text": text})
            if missing:
                unresolved.append({"cell": key, "desc": desc, "deps": missing,
                                   "kind": "复合字段部分来源缺失（该格已按已有字段生成，"
                                           "请向上游提取方补齐后重跑）"})
            if col and spec.get("citation") != "none":
                for vdep, pv in provs:
                    col.add(pv, vdep)
            continue
        dep = spec.get("data_dep")
        if not dep:
            cells.append({"row": row, "col": col_i, "text": "【待填写：%s】" % (desc or key)})
            unresolved.append({"cell": key, "desc": desc, "kind": "无data_dep（主观/待定）"})
            continue
        found, val, prov = resolve_prov(data, dep)
        if not found:
            # 路径不存在：数组越界（如首发无第2项目公司）→ 跳过该格；否则占位
            if "[1]" in dep or "[2]" in dep:
                continue
            cells.append({"row": row, "col": col_i, "text": "【待填写：%s】" % (desc or dep)})
            unresolved.append({"cell": key, "dep": dep, "kind": "extracted_data中无此字段"})
        elif is_blank(val) or isinstance(val, (dict, list)):
            cells.append({"row": row, "col": col_i, "text": "【待填写：%s】" % (desc or dep)})
            unresolved.append({"cell": key, "dep": dep, "kind": "字段为空/结构不符"})
        else:
            cells.append({"row": row, "col": col_i, "text": fmt_by(val, spec.get("format"))})
            if col and spec.get("citation") != "none":
                col.add(prov, dep)

    item = {"locate": locate, "cells": cells}
    if hint.get("clean_headers") or target.get("clean_headers"):
        item["clean_headers"] = True
    # 行结构调整（fill_docx 内部固定顺序：delete_rows → insert_rows → append_rows → cells，
    # 所以 rows_map 的 R{r}C{c} 一律按**结构调整后的最终表**编号）
    dr = target.get("delete_rows") or hint.get("delete_rows")
    if dr:
        item["delete_rows"] = [{"row": d["row"]} for d in dr
                               if isinstance(d, dict) and "row" in d and cond_ok(data, d)]
        if not item["delete_rows"]:
            item.pop("delete_rows")
    for key in ("insert_rows", "append_rows"):
        ops = row_ops_from_target(target, data, key)
        if ops:
            item[key] = ops
    if hwarn:
        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "table_header_unresolved", "table_index": target.get("table_index"),
                     "cells": hwarn,
                     "reason": "表头改写取值失败，模版占位表头（如「第n-3年」）将原样残留 —— "
                               "补齐对应字段后重跑本 phase，或由主agent手工改这几格"})
    if unresolved:
        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "table_unresolved", "cells": unresolved,
                     "reason": "以下单元格已用占位符，主agent确认是否能从材料回补"})
    if col:
        note = col.note_item()
        if note is not None and notes is not None:
            notes.append(note)
        td = col.todo_item()
        if td:
            todo.append(td)
    return item



def build_stage6_entry(entry, data, todo, citation_mode="inline"):
    """生成需延迟到 phase6（最后一批）应用的「整表重建」。

    为什么必须重建而不能原地改（实测模版结构）：
      - 表16：模版是 21行×3列的「行业↔手续」对照表，实际产出要求 7 列（按表15格式列本项目手续），
        python-docx 无法改变已有表格的列数；
      - 表15：模版在 序号/手续名称 两列按 9 大类做了纵向合并（另有 1 行 gridSpan=5），
        合并块行数写死为模版行数，与实际手续行数（实测 33+ 行）对不上，原地插行必然错位；
      - 表22：模版在 阶段 列按阶段纵向合并，实际还需「交易环节跨列小标题行+阶段小计行+合计行」。
    统一做法：insert_tables 新建一张干净表（可自由合并）+ delete_table 删除原表。
    删表会使其后所有 table_index 减1，因此必须放在最后一批。

    来源标注：重建表都是 7~9 列的宽表，只支持 remark_col（蓝图须给 citation_col，
    通常指「备注」列）。不支持 table_note——新表由 insert_tables 用段落锚点插入，
    此时表下注也锚同一段落，会插到表之前，位置错误。
    """
    target = entry.get("target", {})
    maps = entry.get("maps", {}) or {}
    warn = []
    rows = flatten_rows(target, data, maps, warn)
    header = target.get("new_table_header") or []
    cols = int(target.get("cols", len(header) or 7))
    anchor = target.get("after_paragraph", "")
    ti = target.get("table_index")
    out = {"tables": [], "insert_tables": []}

    cit_col = None
    if citation_mode != "none" and str(target.get("citation_placement") or "remark_col") != "none":
        raw = target.get("citation_col")
        if raw is not None:
            try:
                cit_col = int(raw)
            except (TypeError, ValueError):
                cit_col = None
        if cit_col is None:
            todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                         "type": "missing_citation_col", "table_index": ti,
                         "reason": "整表重建的表未在蓝图指定 citation_col（写入来源括注的列，"
                                   "通常是「备注」列索引），本表将不带来源标注。"
                                   "如需标注请在 phase5_blueprints.json 的 target 补 citation_col"})

    if not rows:
        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "table_rebuild_skipped", "table_index": ti,
                     "reason": "数据源(%s)为空，未生成重建计划，模版表保持原样。"
                               "若本项目确实不涉及，在正文写明「不涉及」；否则向上游提取方补数据"
                               % target.get("data_source"),
                     "detail": warn})
        return out

    cells = [{"row": 0, "col": c, "text": str(t)} for c, t in enumerate(header[:cols])]
    merges = []
    n_cit = 0
    n_missing = 0
    for i, row in enumerate(rows):
        fr = i + 1                      # 第0行为表头
        span = row.get("span")          # [from_col, to_col]：合并行/小计/合计/span_rules 命中行
        rc = row.get("cit")
        is_group_row = bool(span and span[0] == 0 and span[1] >= cols - 1)
        if rc is not None and rc.get("type") == "pending":
            rc = None
            if not is_group_row:
                n_missing += 1
        for c, val in enumerate(row["values"][:cols]):
            if span and span[0] < c <= span[1]:
                continue                # 跨列区间只写起始列
            attach = (rc if (cit_col is not None and c == cit_col and rc) else None)
            if val == "" and attach is None:
                continue                # 跨列区间只写起始列
            cell = {"row": fr, "col": c, "text": val}
            if attach:
                cell["citation"] = attach
                n_cit += 1
            cells.append(cell)
        if span and span[1] > span[0]:
            base = row["values"][span[0]] if span[0] < len(row["values"]) else ""
            merges.append({"row": fr, "from_col": span[0], "to_col": min(span[1], cols - 1),
                           "text": row.get("merge_text") or base})

    out["insert_tables"].append({
        "after_paragraph": anchor,
        "rows": len(rows) + 1,
        "cols": cols,
        "style": target.get("style", "Table Grid"),
        "cells": cells,
        "merge_cells": merges,
    })
    out["tables"].append({"locate": {"table_index": ti}, "delete_table": True})
    n_ph = sum(1 for c in cells if "【待填写" in c["text"])
    if n_missing:
        todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                     "type": "missing_citation", "table_index": ti,
                     "fields": ["%d 行缺 attachment_no/_attachment_no" % n_missing],
                     "reason": "重建表中有 %d 行取不到材料编号，这些行不带来源标注。"
                               "第五章合规记录的 attachment_no 是模版硬要求，请向上游提取方补录" % n_missing})
    todo.append({"no": entry.get("no"), "chapter": entry.get("chapter"),
                 "type": "table_rebuild_deferred", "table_index": ti,
                 "reason": "已生成整表重建到 .phase6.json：新建 %d行×%d列（含%d处跨列合并）+ 删除原表%s。"
                           "必须最后一批应用（删表使其后 table_index 减1）"
                           % (len(rows) + 1, cols, len(merges), ti),
                 "generated_rows": len(rows), "placeholders": n_ph,
                 "citations": n_cit, "citation_col": cit_col, "warnings": warn})
    return out



def gen_base_vars(data):
    """从 extracted_data 生成 base_vars（尽力映射，缺的填null）。

    同时生成 citations 段：每个全局变量对应的**现成 citation 对象**。
    这是防子agent编造附件编号/页码的关键——子agent写正文时直接抄
    base_vars.citations.<变量名> 即可，不必自己去 extracted_data 里翻溯源字段。
    """
    cits = {}

    def g(path, key=None, fmt=None):
        found, v, prov = resolve_prov(data, path)
        if not (found and not is_blank(v)):
            return None
        if key:
            c = make_citation(prov, path)
            if c.get("type") != "pending":
                cits[key] = c
        # base_vars_template 承诺「金额字段已格式化为千分位+2位小数的字符串」
        return fmt_by(v, fmt) if fmt else v
    subs = []
    found, sp = resolve(data, "sub_projects")
    if found and isinstance(sp, list):
        for i, s in enumerate(sp):
            if not isinstance(s, dict):
                continue
            # 第一章表1 的口径全部进 base_vars——后续各章引用资产规模/面积/机柜/权属年限时
            # 必须取这里的值，不得另行推定或换用别处的数（避免"表1一个数、正文另一个数"）
            subs.append({k: s.get(k) for k in (
                "name", "project_company", "location", "asset_scope", "underlying_asset",
                "land_area", "building_area", "billing_rack_count", "total_rack_power",
                "construction_cost", "construction_start", "construction_completion",
                "operation_start", "property_right_start", "property_right_end",
                "remaining_years")})
            subs[-1]["construction_cost"] = (fmt_amount(s["construction_cost"])
                                             if not is_blank(s.get("construction_cost")) else None)
            c = make_citation(_prov_of(s), "sub_projects[%d]" % i)
            if c.get("type") != "pending":
                cits["sub_projects[%d]" % i] = c
    orig = []
    orig_short = None
    found, og = resolve(data, "entities.originators")
    if found and isinstance(og, list):
        orig = [o.get("name") for o in og if o.get("name")]
        for i, o in enumerate(og):
            if not isinstance(o, dict):
                continue
            if orig_short is None:
                orig_short = o.get("short_name")
            c = make_citation(_prov_of(o), "entities.originators[%d]" % i)
            if c.get("type") != "pending":
                cits["originators[%d]" % i] = c
    # 运营管理机构（第六章用）：可能有多家，逐家点名是（三）1 与（四）引导段的硬要求
    op_names, op_short = [], None
    found, oms = resolve(data, "entities.operation_managers")
    if found and isinstance(oms, list):
        for i, om in enumerate(oms):
            if not isinstance(om, dict):
                continue
            if om.get("name"):
                op_names.append(om["name"])
            if op_short is None:
                op_short = om.get("short_name") or om.get("name")
            c = make_citation(_prov_of(om), "entities.operation_managers[%d]" % i)
            if c.get("type") != "pending":
                cits["operation_managers[%d]" % i] = c
    bv = {
        "project_name": g("project_info.project_name", "project_name"),
        "fund_name": g("project_info.fund_name", "fund_name"),
        "project_type": g("project_info.project_type", "project_type"),
        "industry": g("project_info.industry", "industry"),
        "issuance_type": g("project_info.issuance_type", "issuance_type"),
        "declaration_base_date": g("project_info.declaration_base_date", "declaration_base_date"),
        "evaluation_value": g("evaluation.total_value", "evaluation_value", "amount"),
        "evaluation_net_value": g("evaluation.total_net_value", "evaluation_net_value", "amount"),
        "total_fund_size": g("project_info.total_fund_size", "total_fund_size", "amount"),
        "exchange": g("project_info.exchange", "exchange"),
        "sub_projects": subs,
        "originators": orig,
        # ---- 第一章专用（表1 与（三）可扩募资产的口径，全文引用以此为准）----
        "originator_short": orig_short or g("expandable_assets.originator_profile.short_name"),
        "expandable_assets_count": g("expandable_assets.summary.total_count"),
        "expandable_assets_scale": g("expandable_assets.summary.total_scale"),
        "expandable_assets_multiple": g("expandable_assets.summary.multiple_of_target",
                                        "expandable_assets_multiple"),
        "fund_manager_name": g("entities.fund_manager.name", "fund_manager_name"),
        "abs_manager_name": g("entities.abs_manager.name", "abs_manager_name"),
        "operation_manager_name": g("entities.operation_manager.name", "operation_manager_name"),
        "law_firm_name": g("entities.law_firm.name", "law_firm_name"),
        "accounting_firm_name": g("entities.accounting_firm.name", "accounting_firm_name"),
        "valuation_agency_name": g("entities.valuation_agency.name", "valuation_agency_name"),
        "tax_advisor_name": g("entities.tax_advisor.name", "tax_advisor_name"),
        "financial_advisor_name": g("entities.financial_advisor.name", "financial_advisor_name"),
        # ---- 第六章「运营管理安排」专用（缺这些变量时该章只能写占位符）----
        "operation_manager_names": "、".join(op_names) if op_names else None,
        "operation_manager_short": op_short,
        "originator_subscription_ratio": g("project_info.originator_subscription_ratio",
                                           "originator_subscription_ratio"),
        "service_agreement_name": g("operation_management.service_agreement.name",
                                    "service_agreement_name"),
        "fee_tier_count": g("operation_management.fee_structure.tier_count", "fee_tier_count"),
        "fee_kpi_metric": g("operation_management.fee_structure.kpi_metric", "fee_kpi_metric"),
        "conflict_commitment_doc_name": g("operation_management.conflict_commitments.doc_name",
                                          "conflict_commitment_doc_name"),
        "conflict_commitment_attachment_no": g(
            "operation_management.conflict_commitments.attachment_no",
            "conflict_commitment_attachment_no"),
    }
    # 只有一家运营管理机构（未填 operation_managers 数组）时，回退到单对象
    # ——见 data_crossref.json 的等价规则 "operation_managers[0]=operation_manager"
    if not bv["operation_manager_names"]:
        bv["operation_manager_names"] = bv.get("operation_manager_name")
    if not bv["operation_manager_short"]:
        bv["operation_manager_short"] = bv.get("operation_manager_name")
    bv["citations"] = cits
    bv["$citations_usage"] = (
        "撰写正文引用上述任一变量时，直接把 citations.<变量名> 作为该 paragraphs 条目的 "
        "citations 数组元素（可加 anchor 指定括注插在正文哪个子串之后）。"
        "citations 中没有的变量说明 extracted_data 缺溯源字段，"
        "写正文时用 {\"type\":\"pending\",\"field\":\"<字段路径>\"}，"
        "严禁自行编造附件编号或页码。话术规范见 templates/citation_rules.json")
    return bv



def adapt_plan_to_base(plan, stage6_tables, stage6_inserts, base_docx, todo, official_path=None):
    """--base-docx：把按官方模版坐标生成的计划适配到实际基底（初稿）。
    真实事故：align_table_index.py 产出的 map 此前没有任何脚本消费，蓝图计划按官方 idx
    打到初稿上必然整体错位（实测初稿删了表2 → 官方 idx3~25 全部 -1，预检全数拦下但无修复通路）。
    本函数一处闭环三件事（不给 --base-docx 时完全不执行，行为不变）：
      ① table_index 按指纹对齐重映射（复用 align_table_index）；基底缺表 → 整条转 todo
      ② 目标格已有实质内容（非空且不含【占位）→ 跳过不覆写（「已填→不动」从口头承诺落到机制）
      ③ delete_rows 越界 → 丢弃；无插行/追加时 cells 越界 → 丢弃（基底被业务增删过行），均记 todo
    """
    from docx import Document
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import align_table_index as ali
    if not official_path:
        assets = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets")
        cand = [os.path.join(assets, f) for f in os.listdir(assets)
                if f.endswith(".docx") and not f.startswith("~$")] if os.path.isdir(assets) else []
        official_path = cand[0] if cand else None
    if not official_path:
        print("⚠️ --base-docx：找不到内置官方模版，跳过基底适配")
        return
    mapping, unmatched, _extra = ali.align(ali.table_sigs(official_path), ali.table_sigs(base_docx))
    doc = Document(base_docx)
    stats = {"remap": 0, "entry_missing": 0, "entry_all_filled": 0,
             "cell_filled": 0, "cell_oob": 0, "del_oob": 0}

    def _filled(text):
        t = (text or "").strip()
        return bool(t) and "【" not in t

    kept = []
    for t in plan.get("tables", []):
        loc = t.get("locate", {})
        idx = loc.get("table_index")
        if idx is None:
            kept.append(t)
            continue
        if str(idx) not in mapping:
            stats["entry_missing"] += 1
            todo.append({"type": "table_missing_in_base", "official_table_index": idx,
                         "hint": "基底（初稿）没有这张官方表：需用 insert_tables 按官方结构在正确位置新建后再填，"
                                 "或与业务确认该表是否被有意删除",
                         "entry": t})
            continue
        real = mapping[str(idx)]
        if real != idx:
            stats["remap"] += 1
        loc["table_index"] = real
        tbl = doc.tables[real]
        n_rows, n_cols = len(tbl.rows), len(tbl.columns)
        has_struct = bool(t.get("insert_rows") or t.get("append_rows"))
        if t.get("delete_rows"):
            kept_del = []
            for d in t["delete_rows"]:
                if d.get("row", 0) >= n_rows:
                    stats["del_oob"] += 1
                    todo.append({"type": "delete_row_out_of_range", "table_index": real, "row": d.get("row"),
                                 "hint": "基底该表只有 %d 行（官方更多），删行条目已丢弃" % n_rows})
                else:
                    kept_del.append(d)
            t["delete_rows"] = kept_del
        n_del = len(t.get("delete_rows") or [])
        cells = []
        for c in t.get("cells", []):
            r, k = c.get("row", 0), c.get("col", 0)
            # cells 的 row 按结构调整后的最终表计算；无插行/追加时最终行数 = 现有行数 - 删行数
            if not has_struct and (r >= n_rows - n_del or k >= n_cols):
                stats["cell_oob"] += 1
                todo.append({"type": "cell_out_of_range", "table_index": real,
                             "cell": {"row": r, "col": k}, "text": c.get("text", "")[:80],
                             "hint": "基底表结构(%dr x %dc，删行后%dr)与官方不同，该格已丢弃，需人工核对落位"
                                     % (n_rows, n_cols, n_rows - n_del)})
                continue
            # 已填保护：仅在无任何行结构操作时可靠（行号未位移），目标格有实质内容则不覆写
            if (not has_struct and n_del == 0 and r < n_rows and k < n_cols
                    and _filled(tbl.rows[r].cells[k].text)):
                stats["cell_filled"] += 1
                # 【冲突可见化】只看行首标签列（col 0）：标签列与蓝图不一致是表结构错乱的强
                # 信号（典型：初稿中介机构表标签列错位）；数据列差异多为等价表述（日期/
                # 金额格式），不报，避免噪音淹没真问题。静默跳过会让错乱表原样混过全部迭代。
                existing = tbl.rows[r].cells[k].text.strip()
                planned = (c.get("text") or "").strip()
                if k == 0 and planned and "【" not in planned and planned != existing:
                    todo.append({"type": "cell_conflict_kept_existing", "table_index": real,
                                 "cell": {"row": r, "col": k},
                                 "existing": existing[:60], "planned": planned[:60],
                                 "hint": "基底该格已有内容但与蓝图计划值不同——若属初稿表结构错乱，"
                                         "需人工重排或按官方模版整表重建后再填"})
                continue
            cells.append(c)
        t["cells"] = cells
        if not (cells or t.get("delete_rows") or has_struct
                or t.get("merge_cells") or t.get("header_cells")):
            stats["entry_all_filled"] += 1
            continue
        kept.append(t)
    plan["tables"] = kept

    # 表下注随表走：所属表在基底缺失则一并丢弃
    if plan.get("insert_paragraphs"):
        kept_notes = []
        for ip in plan["insert_paragraphs"]:
            ai = ip.get("after_table_index")
            if ai is None:
                kept_notes.append(ip)
            elif str(ai) in mapping:
                ip["after_table_index"] = mapping[str(ai)]
                kept_notes.append(ip)
            else:
                todo.append({"type": "table_note_dropped", "official_table_index": ai,
                             "hint": "所属表在基底缺失，表下注随表丢弃"})
        plan["insert_paragraphs"] = kept_notes

    # phase6 整表重建条目只做 idx 重映射（重建本身不受基底行差影响）
    def _remap_any(obj):
        if isinstance(obj, dict):
            for key in ("table_index", "after_table_index"):
                if isinstance(obj.get(key), int) and str(obj[key]) in mapping:
                    obj[key] = mapping[str(obj[key])]
            for v in obj.values():
                _remap_any(v)
        elif isinstance(obj, list):
            for v in obj:
                _remap_any(v)
    _remap_any(stage6_tables)
    _remap_any(stage6_inserts)

    print("基底适配(--base-docx): 重映射 %(remap)d 表 | 基底缺表转todo %(entry_missing)d | "
          "整条已填跳过 %(entry_all_filled)d | 已填格跳过 %(cell_filled)d | "
          "越界格丢弃 %(cell_oob)d | 越界删行丢弃 %(del_oob)d" % stats)
    if unmatched:
        print("⚠️ 基底缺官方表 %d 张（详见 todo 的 table_missing_in_base）" % len(unmatched))


def main():
    ap = argparse.ArgumentParser(description="phase0-2/5 fill_plan 确定性生成器")
    ap.add_argument("--blueprint", required=True, help="phaseN_blueprints.json 路径")
    ap.add_argument("--extracted", required=True, help="extracted_data.json 路径")
    ap.add_argument("--output", required=True, help="输出 fill_plan_phaseN.json")
    ap.add_argument("--base-docx", default=None,
                    help="实际填充基底 docx（初稿）。提供时：table_index 自动对齐到基底、"
                         "已填单元格跳过不覆写、越界条目降级进 todo；不提供则按官方模版坐标原样输出")
    ap.add_argument("--base-vars-out", default=None, help="（仅phase0）同时生成 base_vars.json")
    ap.add_argument("--citation-mode", choices=["inline", "none"], default="inline",
                    help="来源标注模式：inline=生成表下注/备注列括注（默认）；none=不生成")
    ap.add_argument("--base-vars", default=None,
                    help="base_vars.json 路径（可选）：auto_paragraphs 的主体权威名单来源，"
                         "用于从 entities 脏数组里选出真主体（强烈建议 phase2 传入）")
    ap.add_argument("--no-placeholder-tables", dest="placeholder_tables",
                    action="store_false", default=True,
                    help="数据源为空的**模版缺失表**不再插骨架表（默认插：表头+一行【待填写】占位）。"
                         "默认插的原因：正文里的表标题段由子agent写死，表不生成就成了"
                         "「有标题没表格」的悬空引用（实测事故：表4-4~表4-15 十二张全悬空）")
    ap.add_argument("--no-compliance-notes", dest="compliance_notes",
                    action="store_false", default=True,
                    help="不生成**合规门槛提示段**（默认生成）。目前只有一项：第四章（三）1.运营时间的"
                         "「运营满3年」原则 —— operated_years 不满 3 年（或读不出）时，"
                         "在正文该小节插一段黄底【待确认：…】提示。**只提示不阻断**，"
                         "任何情况下都不会让本脚本或校验报 FAIL")
    args = ap.parse_args()

    try:
        assert_handoff_ready(args.extracted, args.output)
    except HandoffGateError as exc:
        print("ERROR: 交接硬门禁阻断生成：%s" % exc, file=sys.stderr)
        sys.exit(3)

    with open(args.blueprint, encoding="utf-8") as f:
        bp = json.load(f)
    with open(args.extracted, encoding="utf-8") as f:
        data = json.load(f)
    # fund_manager_profile 兜底：extracted 缺管理人档案时用蓝图同目录的内置权威档案
    # （表9 数据源；档案里的【待填写：…】按占位处理，不算取到真值，validate 会提示）
    prof_path = os.path.join(os.path.dirname(os.path.abspath(args.blueprint)),
                             "fund_manager_profile.json")
    if not data.get("fund_manager_profile") and os.path.exists(prof_path):
        with open(prof_path, encoding="utf-8") as f:
            data["fund_manager_profile"] = json.load(f)
    base_vars = None
    if args.base_vars:
        with open(args.base_vars, encoding="utf-8") as f:
            base_vars = json.load(f)

    todo = []
    tables = []
    notes = []
    paras = []
    rebuilds = []
    # 延迟应用的批次：{stage: {"tables": [...], "insert_tables": [...]}}
    # phase6 = 整表重建（含 delete_table，table_index -1）
    # phase7 = 模版缺失表的新建（只 insert_tables，table_index +1）
    # 两者都会位移 table_index，必须排在所有按 table_index 定位的批次之后
    staged = {}
    for entry in bp.get("entries", []):
        if entry.get("disabled") is True:
            continue
        stage = entry.get("apply_stage")
        if stage in ("phase6", "phase7"):
            raw = entry.get("%s_plan" % stage) or entry.get("phase6_plan") or {}
            built = build_stage6_entry(entry, data, todo, args.citation_mode,
                                       placeholder_empty=args.placeholder_tables)
            slot = staged.setdefault(stage, {"tables": [], "insert_tables": []})
            for t in (built.get("tables") or raw.get("tables") or []):
                slot["tables"].append(t)
            for t in (built.get("insert_tables") or raw.get("insert_tables") or []):
                slot["insert_tables"].append(t)
            continue
        item = gen_entry(entry, data, todo, notes, args.citation_mode, paras, base_vars,
                         compliance_notes=args.compliance_notes, rebuilds=rebuilds)
        if item:
            tables.append(item)

    plan = {"tables": tables, "paragraphs": paras}
    if rebuilds:
        plan["rebuild_tables"] = rebuilds
    # 蓝图声明了所属章（$chapter）且本次有段落条目时透传，
    # 让 fill_docx 预检能拦截 match 误命中其它章节的情况
    if (paras or rebuilds) and bp.get("$chapter"):
        plan["chapter"] = bp["$chapter"]
    if notes:
        plan["insert_paragraphs"] = notes
    if args.base_docx:
        # staged 各批次里存的是同一批 dict 对象，adapt 内部的 idx 重映射原地生效
        _st = [t for s in staged.values() for t in s["tables"]]
        _si = [t for s in staged.values() for t in s["insert_tables"]]
        adapt_plan_to_base(plan, _st, _si, args.base_docx, todo)
        tables = plan["tables"]              # 适配后统计按实际输出算
        notes = plan.get("insert_paragraphs", [])
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(plan, f, ensure_ascii=False, indent=2)

    todo_path = args.output.replace(".json", ".todo.json")
    with open(todo_path, "w", encoding="utf-8") as f:
        json.dump(todo, f, ensure_ascii=False, indent=2)

    staged_paths = []
    for stage in ("phase6", "phase7"):
        slot = staged.get(stage)
        if not slot or not (slot["tables"] or slot["insert_tables"]):
            continue
        p = args.output.replace(".json", ".%s.json" % stage)
        with open(p, "w", encoding="utf-8") as f:
            json.dump({"paragraphs": [], "insert_tables": slot["insert_tables"],
                       "tables": slot["tables"]}, f, ensure_ascii=False, indent=2)
        staged_paths.append((stage, p, len(slot["insert_tables"]), len(slot["tables"])))

    # 蓝图声明 $emit_empty_stages：该批次已整体迁到 rebuild_tables（主 plan 直出），
    # 仍写一个空 .phaseN.json——旧流程文档/脚本按固定文件名找产物，缺文件会误判漏跑
    for stage in (bp.get("$emit_empty_stages") or []):
        if any(s == stage for s, _p, _a, _b in staged_paths):
            continue
        p = args.output.replace(".json", ".%s.json" % stage)
        with open(p, "w", encoding="utf-8") as f:
            json.dump({"paragraphs": [], "insert_tables": [], "tables": []},
                      f, ensure_ascii=False, indent=2)
        print("兼容旧流程：%s 批次已迁移至 rebuild_tables（主 plan 直出），已写空产物 %s"
              % (stage, p))

    n_cells = sum(len(t.get("cells", [])) for t in tables)
    n_ph = sum(1 for t in tables for c in t.get("cells", []) if "【待填写" in c["text"])
    n_struct = sum(len(t.get("insert_rows", [])) + len(t.get("append_rows", []))
                   + len(t.get("delete_rows", [])) + len(t.get("merge_cells", [])) for t in tables)
    n_cell_cit = sum(1 for t in tables for c in t.get("cells", []) if c.get("citation"))
    n_note_cit = sum(len(n.get("citations", [])) for n in notes)
    n_table_notes = sum(1 for n in notes if n.get("citations"))
    n_miss = sum(1 for t in todo if t.get("type") in ("missing_citation", "missing_citation_col"))
    print("=== fill_plan 生成完成: %s ===" % args.output)
    print("表格条目: %d   单元格: %d（其中占位符 %d）   行结构操作: %d"
          % (len(tables), n_cells, n_ph, n_struct))
    if rebuilds:
        n_rb_new = sum(1 for r in rebuilds
                       if not (r.get("locate") or {}).get("title_keyword"))
        print("表格重建条目: %d（其中纯新建 %d 张走 create_after 锚）—— caption 一律「表#」，"
              "交付前跑 renumber_tables.py 统一编号" % (len(rebuilds), n_rb_new))
    if paras:
        n_para_ph = sum(1 for p in paras if "【待填写" in p.get("replace", ""))
        print("模板段直出: %d 处替换（其中 %d 处含【待填写】，见 todo 的 missing_template_fields）"
              % (len(paras), n_para_ph))
    if args.citation_mode == "none":
        print("来源标注: 已用 --citation-mode none 关闭")
    else:
        print("来源标注: 备注列括注 %d 处 + 表下注 %d 段（含 %d 条来源）"
              % (n_cell_cit, n_table_notes, n_note_cit))
        if n_miss:
            print("⚠️ 有 %d 项缺溯源，见 todo 的 missing_citation/missing_citation_col —— "
                  "这些内容将不带来源标注，请向上游提取方为对应字段补 _attachment_no/_doc_name/_page"
                  % n_miss)
    print("待主agent处理条目（段落/公式/未解析/缺溯源）: %d → %s" % (len(todo), todo_path))
    # ---- 表格数据缺口的显式播报（第四章 15 表事故的直接对策）----
    ph_tables = [t for t in todo if t.get("type") == "table_new_placeholder"]
    skipped = [t for t in todo if t.get("type") in ("table_new_skipped", "table_rebuild_skipped")]
    if ph_tables:
        print("⚠️ 骨架表 %d 张（数据源为空，只插了表头+一行【待填写】占位，**不是交付形态**）："
              % len(ph_tables))
        for t in ph_tables:
            print("    %s ← %s" % (t.get("table_name"), t.get("data_source")))
        print("   → 必须向用户/上游列明这些表缺数据（第四章的 valuation_params.* 缺失，"
              "典型根因是评估报告只提取了前几页摘要）；补录后重跑本 phase 并重新应用")
    if skipped:
        print("⚠️ 未生成的表 %d 张（数据源为空且未插骨架表）：%s —— "
              "正文若已写表标题/「下表列示…」即为悬空引用，必须补数据或删引用"
              % (len(skipped), "、".join(str(t.get("table_name")) for t in skipped)))
    # ---- 合规门槛提示（运营满3年）：只提示不阻断，但必须在终端明确播报 ----
    op_notes = [t for t in todo
                if t.get("type") in ("operation_years_short", "operation_years_unverified")]
    for t in op_notes:
        if t.get("type") == "operation_years_short":
            print("⚠️ 【运营满3年门槛】operated_years=%s，**不满 %g 年** —— 基础设施项目原则上"
                  "要求已运营满 %g 年。%s"
                  % (t.get("operated_years"), t.get("min_years"), t.get("min_years"),
                     ("已在正文（三）1.运营时间的锚点「%s」后插入黄底【待确认】提示段"
                      % t.get("anchor")) if t.get("note_inserted") else "未插提示段"))
            print("   → %s不满 %g 年时必须在本节写明「能够实现长期稳定收益的原因」；"
                  "核实后请把提示段删除，不要留到交付稿。"
                  "**本项只提示、不阻断，不会让任何脚本报 FAIL**"
                  % ("extracted_data 已提供 under_3_years_reason，请确认已写进正文；"
                     if t.get("has_reason") else
                     "extracted_data 的 under_3_years_reason 为空 —— ", t.get("min_years")))
        else:
            print("ℹ️ 【运营满3年门槛】无法核验（operated_years %s）—— %s"
                  % ("为空" if t.get("kind") == "missing"
                     else "读不出年限：%r" % t.get("operated_years"),
                     "已插入黄底【待确认】提示段" if t.get("note_inserted") else "仅记录未插段"))
    for stage, p, n_ins, n_del in staged_paths:
        print("⚠️ 需延后应用的批次[%s]: %s（新建表 %d 张%s）"
              % (stage, p, n_ins, "，删除原表 %d 张" % n_del if n_del else ""))
    if staged_paths:
        print("   应用顺序：各章 fill_plan（按 table_index 定位的都要排在前面）→ 附件1/2 清理"
              " → .phase6.json（整表重建，删表使 table_index -1）"
              " → .phase7.json（模版缺失表新建，插表使 table_index +1，**最后**）")

    if args.base_vars_out:
        bv = gen_base_vars(data)
        with open(args.base_vars_out, "w", encoding="utf-8") as f:
            json.dump(bv, f, ensure_ascii=False, indent=2)
        n_null = sum(1 for k, v in bv.items()
                     if not k.startswith("$") and k != "citations" and v in (None, [], {}))
        print("base_vars.json 已生成（%d 个字段为null，%d 个变量带现成来源标注）: %s"
              % (n_null, len(bv.get("citations", {})), args.base_vars_out))
        if not bv.get("citations"):
            print("⚠️ base_vars.citations 为空：extracted_data 完全没有 _attachment_no/_source，"
                  "子agent将无法标注来源。请向上游提取方补溯源字段")

    # 生成计划也有硬门禁：整表空缺或有值无来源时，计划文件保留供定位问题，
    # 但 exit=4 阻止后续 Word 写入。普通单元格占位仍允许，最终由成品校验统一阻断。
    if ph_tables or skipped or n_miss:
        print("❌ PLAN BLOCKED：整表占位=%d、整表跳过=%d、缺来源=%d。"
              "先补 extracted_data 后重跑；当前 plan 仅供诊断，不得写入Word。"
              % (len(ph_tables), len(skipped), n_miss), file=sys.stderr)
        sys.exit(4)



if __name__ == "__main__":
    main()
